import sqlite3
import csv
import io
import os
import base64
import binascii
import hashlib
import uuid
from dotenv import load_dotenv

load_dotenv()
import json
from functools import wraps
from datetime import datetime, date, time, timedelta, timezone
from flask import Flask, render_template, request, redirect, url_for, flash, g, session, jsonify, Response, send_file, abort
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash

try:
    from authlib.integrations.flask_client import OAuth
    from authlib.integrations.base_client.errors import OAuthError
except ImportError:  # Keeps the setup page usable until requirements are installed.
    OAuth = None
    OAuthError = Exception

from invoice_pdf import build_invoice_pdf

DB_PATH = "leads.db"
ROLES = ["telecaller", "field_sales", "manager"]
# QuantumLoop operates on India Standard Time. A fixed offset avoids depending on
# the optional Windows IANA timezone database while remaining exact (IST has no DST).
APP_TIMEZONE = timezone(timedelta(hours=5, minutes=30), name="Asia/Kolkata")
FIELD_WORK_AUTO_END = time(18, 30)
NO_CARD_REASONS = ["No card available", "Forgot to ask", "Contact declined", "Other"]
VISIT_OUTCOMES = ["Interested", "Follow-up Required", "Proposal Requested", "Not Interested", "No Decision", "Contact Unavailable", "Other"]

# ---------- New pipeline ----------
STAGES = ["New", "Attempted", "Contacted", "Follow Up", "Interested",
          "Qualified", "Proposal Sent", "Negotiation", "Won", "Lost"]
# Stages a manager can see across all owners (escalated / advanced pipeline)
ESCALATE_STAGES = {"Interested", "Qualified", "Proposal Sent", "Negotiation", "Won", "Lost"}
# Map legacy 6-stage values to the new pipeline (run once as a migration)
LEGACY_STAGE_MAP = {
    "Not Interested": "Lost",
    "Converted": "Won",
}

LOST_REASONS = ["No Budget", "Already Has Vendor", "No Requirement", "Wrong Contact",
                "No Response", "Competitor Won", "Other"]

SERVICES = [
    "Website Design", "Website Redesign", "Ecommerce", "SEO", "Google Ads",
    "Social Media Marketing", "Branding", "Graphic Designing", "ERP", "CRM",
    "Mobile App", "Custom Software", "Hardware Supply", "Cyber Security",
    "Server Services", "Networking", "Structured Cabling", "Cloud Services",
    "CCTV & Surveillance", "AMC", "Other",
]

WEBSITE_STATUS = ["No Website", "Has Website", "Under Development"]
WEBSITE_QUALITY = ["Excellent", "Good", "Average", "Poor", "Broken"]
GOOGLE_PRESENCE = ["Excellent", "Average", "Poor", "Not Found"]
SOCIAL_PRESENCE = ["Active", "Moderate", "Poor", "None"]
DECISION_MAKER_OPTIONS = ["Yes", "No", "Unknown"]
INTEREST_LEVELS = ["Hot", "Warm", "Cold"]
PRIORITIES = ["High", "Medium", "Low"]
BUDGET_RANGES = ["Under 25k", "25k-50k", "50k-1L", "1L-2L", "2L-5L", "5L+"]
COMPANY_STATUSES = ["Prospect", "Customer", "Inactive"]

app = Flask(__name__)
app.secret_key = os.environ["SECRET_KEY"]
app.config.update(
    SESSION_COOKIE_SECURE=os.environ.get("SESSION_COOKIE_SECURE", "true").strip().lower() in {"1", "true", "yes", "on"},
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(
        days=max(1, int(os.environ.get("LOGIN_SESSION_DAYS", "3650")))
    ),
    SESSION_REFRESH_EACH_REQUEST=True,
    MAX_CONTENT_LENGTH=12 * 1024 * 1024,
)
# Only trust forwarded host/protocol headers when the deployment has a known proxy.
if os.environ.get("TRUST_PROXY_HEADERS", "false").strip().lower() in {"1", "true", "yes", "on"}:
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

oauth = OAuth(app) if OAuth else None
google = None
if oauth and os.environ.get("GOOGLE_CLIENT_ID") and os.environ.get("GOOGLE_CLIENT_SECRET"):
    google = oauth.register(
        name="google",
        client_id=os.environ["GOOGLE_CLIENT_ID"],
        client_secret=os.environ["GOOGLE_CLIENT_SECRET"],
        server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
        client_kwargs={"scope": "openid email profile"},
    )


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def _add_cols(db, table, coldefs):
    """coldefs: list of (name, sql_type_and_default)"""
    existing = [r[1] for r in db.execute(f"PRAGMA table_info({table})")]
    for name, decl in coldefs:
        if name not in existing:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {decl}")


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            name TEXT NOT NULL,
            role TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
        """
    )
    _add_cols(db, "users", [
        ("email", "TEXT"),
        ("google_sub", "TEXT"),
        ("last_login_at", "TEXT"),
    ])
    # Older installations sometimes used an email address as the username.
    db.execute(
        "UPDATE users SET email=LOWER(TRIM(username)) "
        "WHERE (email IS NULL OR TRIM(email)='') AND username LIKE '%@%'"
    )
    db.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_users_email_nocase "
        "ON users(email COLLATE NOCASE) WHERE email IS NOT NULL AND TRIM(email) <> ''"
    )
    db.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_users_google_sub "
        "ON users(google_sub) WHERE google_sub IS NOT NULL AND TRIM(google_sub) <> ''"
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS admin_login_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            remote_addr TEXT NOT NULL,
            success INTEGER NOT NULL DEFAULT 0,
            attempted_at TEXT NOT NULL
        )
        """
    )
    db.execute(
        "CREATE INDEX IF NOT EXISTS idx_admin_login_attempts_lookup "
        "ON admin_login_attempts(username, remote_addr, attempted_at)"
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS leads (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            phone TEXT NOT NULL,
            email TEXT,
            source TEXT,
            stage TEXT NOT NULL DEFAULT 'New',
            tags TEXT,
            follow_up_date TEXT,
            owner_id INTEGER,
            created_at TEXT NOT NULL
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS call_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lead_id INTEGER NOT NULL,
            outcome TEXT,
            note TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (lead_id) REFERENCES leads (id) ON DELETE CASCADE
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS companies (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            industry TEXT,
            phone TEXT,
            email TEXT,
            website TEXT,
            notes TEXT,
            created_at TEXT NOT NULL
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS contacts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            company_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            designation TEXT,
            phone TEXT,
            email TEXT,
            notes TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (company_id) REFERENCES companies (id) ON DELETE CASCADE
        )
        """
    )
    # Invoice module migrations and tables
    company_cols = [r[1] for r in db.execute("PRAGMA table_info(companies)")]
    for col in ("billing_address", "gstin", "pan"):
        if col not in company_cols:
            db.execute(f"ALTER TABLE companies ADD COLUMN {col} TEXT")

    db.execute("""
        CREATE TABLE IF NOT EXISTS company_settings (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            name TEXT, address_lines TEXT, website TEXT, gstin TEXT, pan TEXT,
            bank_name TEXT, ifsc TEXT, account_no TEXT, bank_branch TEXT, upi_id TEXT,
            default_terms TEXT, default_notes TEXT, invoice_prefix TEXT DEFAULT 'QL'
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT, invoice_no TEXT NOT NULL UNIQUE, invoice_date TEXT NOT NULL,
            company_id INTEGER, bill_name TEXT NOT NULL, bill_address TEXT, bill_gstin TEXT, bill_pan TEXT, bill_website TEXT,
            ship_name TEXT, ship_address TEXT, ship_gstin TEXT, ship_pan TEXT, ship_website TEXT,
            notes TEXT, terms TEXT, cgst_pct REAL DEFAULT 9, sgst_pct REAL DEFAULT 9, received_amount REAL DEFAULT 0,
            total_amount REAL DEFAULT 0, created_by INTEGER, created_at TEXT NOT NULL
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS invoice_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT, invoice_id INTEGER NOT NULL, item_name TEXT NOT NULL, code TEXT,
            qty REAL NOT NULL, rate REAL NOT NULL, tax_pct REAL NOT NULL,
            FOREIGN KEY (invoice_id) REFERENCES invoices(id) ON DELETE CASCADE
        )
    """)
    if not db.execute("SELECT 1 FROM company_settings WHERE id=1").fetchone():
        db.execute("""INSERT INTO company_settings
            (id,name,address_lines,website,gstin,pan,bank_name,ifsc,account_no,bank_branch,upi_id,default_terms,default_notes,invoice_prefix)
            VALUES (1,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
            'QUANTUMLOOP PRIVATE LIMITED',
            'H.No: 1-4-214/26P, P NO26, Sangamitra Enclave, Sainikpuri Road, Hyderabad, Telangana, 500062',
            'www.quantumloop.online', '36AABCQ2756M1Z2', 'AABCQ2756M', 'QuantumLoop Private Limited',
            'IDIB000A135', '8316619617', 'As Rao Nagar', 'quantumloop@okhdfcbank',
            'Quotation valid for 30 days from issue date.\n50% advance required to initiate development.\nDelivery timelines commence upon receipt of advance payment.\nInfrastructure subscriptions are billed annually.\nAdditional features not listed in this quotation will be charged separately.\nGoods/Services once sold will not be taken back.\nSubject to Hyderabad jurisdiction only.',
            '', 'QL'))

    # migrate: add owner_id to leads if this db predates it
    cols = [r[1] for r in db.execute("PRAGMA table_info(leads)")]
    if "owner_id" not in cols:
        db.execute("ALTER TABLE leads ADD COLUMN owner_id INTEGER")

    # ---------- New CRM: Company becomes the center of everything ----------

    _add_cols(db, "companies", [
        ("alternate_phone", "TEXT"),
        ("address", "TEXT"),
        ("city", "TEXT"),
        ("state", "TEXT"),
        ("pincode", "TEXT"),
        ("country", "TEXT DEFAULT 'India'"),
        ("cin", "TEXT"),
        ("status", "TEXT DEFAULT 'Prospect'"),
        ("account_owner_id", "INTEGER"),
        ("updated_at", "TEXT"),
        ("gbp_link", "TEXT"),
    ])

    _add_cols(db, "contacts", [
        ("alternate_phone", "TEXT"),
        ("linkedin", "TEXT"),
        ("is_primary", "INTEGER DEFAULT 0"),
        ("updated_at", "TEXT"),
    ])

    _add_cols(db, "leads", [
        ("company_id", "INTEGER"),
        ("contact_id", "INTEGER"),
        ("industry", "TEXT"),
        ("designation", "TEXT"),
        ("alternate_phone", "TEXT"),
        ("website_status", "TEXT"),
        ("website_url", "TEXT"),
        ("website_quality", "TEXT"),
        ("google_presence", "TEXT"),
        ("social_presence", "TEXT"),
        ("services_required", "TEXT"),
        ("other_service_description", "TEXT"),
        ("decision_maker", "TEXT"),
        ("interest_level", "TEXT"),
        ("priority", "TEXT"),
        ("budget_range", "TEXT"),
        ("estimated_project_value", "TEXT"),
        ("expected_timeline", "TEXT"),
        ("lost_reason", "TEXT"),
        ("qualification_notes", "TEXT"),
        ("requirement_notes", "TEXT"),
        ("internal_notes", "TEXT"),
        ("updated_at", "TEXT"),
    ])

    _add_cols(db, "call_logs", [
        ("user_id", "INTEGER"),
    ])

    _add_cols(db, "invoices", [
        ("bill_contact_name", "TEXT"),
        ("bill_contact_email", "TEXT"),
        ("bill_contact_phone", "TEXT"),
        ("ship_state", "TEXT"),
        ("tax_mode", "TEXT DEFAULT 'INTRA'"),
        ("igst_pct", "REAL DEFAULT 0"),
    ])

    # Migrate legacy stage names -> new pipeline (idempotent; no-op after first run)
    for old_stage, new_stage in LEGACY_STAGE_MAP.items():
        db.execute("UPDATE leads SET stage = ? WHERE stage = ?", (new_stage, old_stage))
        if old_stage == "Not Interested":
            db.execute(
                "UPDATE leads SET lost_reason = COALESCE(lost_reason, 'Other') WHERE stage = 'Lost' AND (lost_reason IS NULL OR lost_reason = '')"
            )

    # Future-proof tables (section 12 of the spec) — created now even if lightly used today.
    db.execute("""
        CREATE TABLE IF NOT EXISTS activity_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            user_name TEXT,
            action TEXT NOT NULL,
            entity_type TEXT NOT NULL,
            entity_id INTEGER,
            entity_label TEXT,
            details TEXT,
            created_at TEXT NOT NULL
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS lead_activities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lead_id INTEGER NOT NULL,
            event TEXT NOT NULL,
            detail TEXT,
            user_id INTEGER,
            user_name TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (lead_id) REFERENCES leads (id) ON DELETE CASCADE
        )
    """)

    # Access grants: gives a user (telecaller or manager) view+edit access to a lead
    # they don't own, WITHOUT changing ownership or copying the record. Ownership
    # (owner_id on leads) never changes via this table.
    db.execute("""
        CREATE TABLE IF NOT EXISTS lead_shares (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lead_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            granted_by INTEGER,
            created_at TEXT NOT NULL,
            UNIQUE(lead_id, user_id),
            FOREIGN KEY (lead_id) REFERENCES leads (id) ON DELETE CASCADE,
            FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE CASCADE
        )
    """)

    # Field Sales daily work sessions and visits. Visit assignment is deliberately
    # separate from lead ownership: telecallers/managers keep CRM responsibility
    # while an executive receives access to the visit and its related record.
    db.execute("""
        CREATE TABLE IF NOT EXISTS field_work_sessions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            work_date TEXT NOT NULL,
            started_at TEXT NOT NULL,
            ended_at TEXT,
            end_source TEXT,
            status TEXT NOT NULL DEFAULT 'active',
            created_at TEXT NOT NULL,
            UNIQUE(user_id, work_date),
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE SET NULL
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS field_visits (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lead_id INTEGER,
            company_id INTEGER,
            assigned_to INTEGER,
            created_by INTEGER,
            origin TEXT NOT NULL,
            scheduled_date TEXT NOT NULL,
            scheduled_time TEXT,
            location_address TEXT NOT NULL,
            purpose TEXT NOT NULL,
            contact_person TEXT,
            contact_designation TEXT,
            contact_phone TEXT,
            contact_email TEXT,
            status TEXT NOT NULL DEFAULT 'scheduled',
            session_id INTEGER,
            checkin_at TEXT,
            checkin_lat REAL,
            checkin_lng REAL,
            checkin_accuracy REAL,
            location_status TEXT,
            completed_at TEXT,
            people_met TEXT,
            discussion TEXT,
            outcome TEXT,
            notes TEXT,
            next_action TEXT,
            follow_up_date TEXT,
            no_card_reason TEXT,
            no_card_explanation TEXT,
            submitted_at TEXT,
            review_status TEXT,
            reviewed_by INTEGER,
            reviewed_at TEXT,
            review_comment TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            FOREIGN KEY (lead_id) REFERENCES leads(id) ON DELETE SET NULL,
            FOREIGN KEY (company_id) REFERENCES companies(id) ON DELETE SET NULL,
            FOREIGN KEY (assigned_to) REFERENCES users(id) ON DELETE SET NULL,
            FOREIGN KEY (created_by) REFERENCES users(id) ON DELETE SET NULL,
            FOREIGN KEY (session_id) REFERENCES field_work_sessions(id) ON DELETE SET NULL,
            FOREIGN KEY (reviewed_by) REFERENCES users(id) ON DELETE SET NULL
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS field_visit_photos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            visit_id INTEGER NOT NULL,
            photo_type TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            mime_type TEXT NOT NULL,
            byte_size INTEGER NOT NULL,
            sha256 TEXT NOT NULL,
            captured_at TEXT NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (visit_id) REFERENCES field_visits(id) ON DELETE CASCADE
        )
    """)
    _add_cols(db, "field_visits", [
        ("contact_designation", "TEXT"),
        ("contact_phone", "TEXT"),
        ("contact_email", "TEXT"),
    ])
    db.execute("CREATE INDEX IF NOT EXISTS idx_field_visits_assignee_date ON field_visits(assigned_to, scheduled_date)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_field_visits_review ON field_visits(review_status)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_field_sessions_user_date ON field_work_sessions(user_id, work_date)")

    db.commit()
    db.close()


def now_str():
    return datetime.now(APP_TIMEZONE).strftime("%Y-%m-%d %H:%M:%S")


def today_str():
    return datetime.now(APP_TIMEZONE).strftime("%Y-%m-%d")


def month_start_str():
    return datetime.now(APP_TIMEZONE).date().replace(day=1).strftime("%Y-%m-%d")


def time_ago(ts):
    if not ts:
        return ""
    try:
        then = datetime.strptime(ts, "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return ts
    delta = datetime.now(APP_TIMEZONE).replace(tzinfo=None) - then
    seconds = delta.total_seconds()
    if seconds < 60:
        return "just now"
    if seconds < 3600:
        return f"{int(seconds // 60)}m ago"
    if seconds < 86400:
        return f"{int(seconds // 3600)}h ago"
    if seconds < 86400 * 30:
        return f"{int(seconds // 86400)}d ago"
    if seconds < 86400 * 365:
        return f"{int(seconds // (86400 * 30))}mo ago"
    return f"{int(seconds // (86400 * 365))}y ago"


# ---------- Activity logging ----------

def log_activity(action, entity_type, entity_id=None, entity_label="", details=""):
    db = get_db()
    user = current_user()
    db.execute(
        "INSERT INTO activity_logs (user_id, user_name, action, entity_type, entity_id, entity_label, details, created_at) VALUES (?,?,?,?,?,?,?,?)",
        (user["id"] if user else None, user["name"] if user else "System", action, entity_type, entity_id, entity_label, details, now_str()),
    )
    db.commit()


def log_lead_event(lead_id, event, detail=""):
    db = get_db()
    user = current_user()
    db.execute(
        "INSERT INTO lead_activities (lead_id, event, detail, user_id, user_name, created_at) VALUES (?,?,?,?,?,?)",
        (lead_id, event, detail, user["id"] if user else None, user["name"] if user else "System", now_str()),
    )
    db.commit()


# ---------- Auth helpers ----------

def current_user():
    if "user_id" not in session:
        return None
    if "user" not in g:
        g.user = get_db().execute("SELECT * FROM users WHERE id = ?", (session["user_id"],)).fetchone()
    return g.user


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not current_user():
            return redirect(url_for("login", next=request.path))
        return f(*args, **kwargs)
    return wrapper


def manager_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = current_user()
        if not user or user["role"] != "manager":
            flash("Only Customer Relations Managers can access that.", "error")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return wrapper


def field_sales_or_manager_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = current_user()
        if not user or user["role"] not in {"manager", "field_sales"}:
            flash("Only managers and Field Sales Executives can access that.", "error")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return wrapper


def _reconcile_field_sessions(db):
    """Close active sessions at 18:30 local time, even if reconciliation runs later."""
    local_now = datetime.now(APP_TIMEZONE)
    rows = db.execute("SELECT * FROM field_work_sessions WHERE status = 'active'").fetchall()
    changed = False
    for row in rows:
        work_day = datetime.strptime(row["work_date"], "%Y-%m-%d").date()
        cutoff = datetime.combine(work_day, FIELD_WORK_AUTO_END).replace(tzinfo=APP_TIMEZONE)
        if local_now >= cutoff:
            db.execute(
                "UPDATE field_work_sessions SET ended_at=?, end_source='system_assumed', status='closed' WHERE id=?",
                (cutoff.strftime("%Y-%m-%d %H:%M:%S"), row["id"]),
            )
            changed = True
    if changed:
        db.commit()


@app.before_request
def require_login():
    open_endpoints = {"login", "setup", "google_login", "google_callback", "manual_admin_login", "static"}
    if request.endpoint in open_endpoints:
        return
    db = get_db()
    _reconcile_field_sessions(db)
    has_users = db.execute("SELECT COUNT(*) c FROM users").fetchone()["c"]
    if not has_users and request.endpoint != "setup":
        return redirect(url_for("setup"))
    if has_users and not current_user():
        return redirect(url_for("login", next=request.path))


def visible_leads_clause(user):
    """Returns (sql_fragment, params) restricting leads to what this user may see.

    Rules:
    - A lead's owner (owner_id) always sees it. Ownership never moves except via
      /leads/<id>/assign for leads that don't have an owner yet.
    - Every telecaller's leads are always visible to every manager, tagged with the
      telecaller's name (managers never "own" a telecaller's lead by seeing it).
    - Anyone explicitly granted access via lead_shares (a "transfer") sees the lead
      too, without it changing owner_id — this is how a manager's own leads, or a
      telecaller's leads, get shared with another telecaller/manager.
    """
    shared_clause = "leads.id IN (SELECT lead_id FROM lead_shares WHERE user_id = ?)"
    visit_clause = "leads.id IN (SELECT lead_id FROM field_visits WHERE assigned_to = ?)"
    if user["role"] == "manager":
        return (
            f"(owner_id = ? OR owner_id IN (SELECT id FROM users WHERE role IN ('telecaller','field_sales')) OR owner_id IS NULL OR {shared_clause})",
            [user["id"], user["id"]],
        )
    if user["role"] == "field_sales":
        return f"(owner_id = ? OR {shared_clause} OR {visit_clause})", [user["id"], user["id"], user["id"]]
    return f"(owner_id = ? OR {shared_clause})", [user["id"], user["id"]]


# ---------- Setup / Auth ----------

def _normalize_email(value):
    return (value or "").strip().lower()


def _valid_email(value):
    email = _normalize_email(value)
    return bool(email and "@" in email and not email.startswith("@") and not email.endswith("@"))


def _safe_next_url(value):
    return value if value and value.startswith("/") and not value.startswith("//") else None


def _bootstrap_emails():
    raw = os.environ.get("GOOGLE_BOOTSTRAP_EMAILS", "")
    return {_normalize_email(value) for value in raw.replace(";", ",").split(",") if value.strip()}


def _google_redirect_uri():
    return os.environ.get("GOOGLE_REDIRECT_URI", "").strip() or url_for("google_callback", _external=True)


def _manual_admin_login_enabled():
    return os.environ.get("MANUAL_ADMIN_LOGIN_ENABLED", "false").strip().lower() in {"1", "true", "yes", "on"}


def _manual_login_remote_addr():
    return (request.remote_addr or "unknown")[:100]


def _manual_login_is_locked(db, username, remote_addr):
    cutoff = (datetime.now(APP_TIMEZONE) - timedelta(minutes=15)).strftime("%Y-%m-%d %H:%M:%S")
    failures = db.execute(
        """SELECT COUNT(*) c FROM admin_login_attempts
           WHERE success=0 AND attempted_at>=? AND (username=? OR remote_addr=?)""",
        (cutoff, username, remote_addr),
    ).fetchone()["c"]
    return failures >= 5


def _record_manual_login_attempt(db, username, remote_addr, success):
    prune_before = (datetime.now(APP_TIMEZONE) - timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S")
    db.execute("DELETE FROM admin_login_attempts WHERE attempted_at<?", (prune_before,))
    db.execute(
        "INSERT INTO admin_login_attempts (username,remote_addr,success,attempted_at) VALUES (?,?,?,?)",
        (username, remote_addr, 1 if success else 0, now_str()),
    )
    if success:
        db.execute(
            "DELETE FROM admin_login_attempts WHERE success=0 AND (username=? OR remote_addr=?)",
            (username, remote_addr),
        )
    db.commit()


@app.route("/setup", methods=["GET", "POST"])
def setup():
    db = get_db()
    has_users = db.execute("SELECT COUNT(*) c FROM users").fetchone()["c"]
    if has_users:
        return redirect(url_for("login"))

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        email = _normalize_email(request.form.get("email"))
        if not name or not _valid_email(email):
            flash("Enter your name and a valid Google account email.", "error")
            return redirect(url_for("setup"))
        db.execute(
            "INSERT INTO users (username, password_hash, name, role, email, created_at) "
            "VALUES (?, 'google-only', ?, 'manager', ?, ?)",
            (email, name, email, now_str()),
        )
        db.commit()
        flash("Manager access created. Continue with Google using that email.", "success")
        return redirect(url_for("login"))

    return render_template("setup.html")


@app.route("/login")
def login():
    if current_user():
        return redirect(url_for("dashboard"))
    next_url = _safe_next_url(request.args.get("next"))
    if next_url:
        session["post_login_next"] = next_url
    return render_template(
        "login.html",
        google_ready=google is not None,
        authlib_ready=OAuth is not None,
        manual_admin_enabled=_manual_admin_login_enabled(),
    )


@app.route("/admin-login", methods=["GET", "POST"])
def manual_admin_login():
    if not _manual_admin_login_enabled():
        abort(404)
    if current_user():
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        db = get_db()
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        remote_addr = _manual_login_remote_addr()

        if _manual_login_is_locked(db, username, remote_addr):
            flash("Too many failed admin login attempts. Try again in 15 minutes.", "error")
            return redirect(url_for("manual_admin_login"))

        user = db.execute(
            "SELECT * FROM users WHERE username=? COLLATE NOCASE AND role='manager'",
            (username,),
        ).fetchone()
        valid_legacy_password = bool(
            user and user["password_hash"] != "google-only"
            and check_password_hash(user["password_hash"], password)
        )

        recovery_username = os.environ.get("MANUAL_ADMIN_USERNAME", "").strip().lower()
        recovery_hash = os.environ.get("MANUAL_ADMIN_PASSWORD_HASH", "").strip()
        valid_recovery_password = bool(
            recovery_username and recovery_hash and username == recovery_username
            and check_password_hash(recovery_hash, password)
        )
        if valid_recovery_password and not user:
            user = db.execute("SELECT * FROM users WHERE role='manager' ORDER BY id LIMIT 1").fetchone()

        if not user or not (valid_legacy_password or valid_recovery_password):
            _record_manual_login_attempt(db, username, remote_addr, False)
            flash("Invalid admin username or password.", "error")
            return redirect(url_for("manual_admin_login"))

        _record_manual_login_attempt(db, username, remote_addr, True)
        next_url = _safe_next_url(session.get("post_login_next")) or url_for("users_list")
        session.clear()
        session.permanent = True
        session["user_id"] = user["id"]
        session["role"] = user["role"]
        g.user = user
        log_activity(
            "Manual admin login",
            "auth",
            user["id"],
            user["name"],
            "Environment recovery credential" if valid_recovery_password else "Legacy manager credential",
        )
        flash("Admin access granted. You can now assign Google emails to users.", "success")
        return redirect(next_url)

    return render_template("manual_admin_login.html")


@app.route("/auth/google")
def google_login():
    if current_user():
        return redirect(url_for("dashboard"))
    if google is None:
        if OAuth is None:
            flash("Google login dependency is not installed. Run pip install -r requirements.txt.", "error")
        else:
            flash("Google login is not configured yet. Add GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET.", "error")
        return redirect(url_for("login"))
    return google.authorize_redirect(_google_redirect_uri())


@app.route("/auth/google/callback")
def google_callback():
    if google is None:
        flash("Google login is not configured.", "error")
        return redirect(url_for("login"))

    try:
        token = google.authorize_access_token()
        userinfo = token.get("userinfo") or google.userinfo(token=token)
    except OAuthError:
        flash("Google sign-in was cancelled or could not be verified. Please try again.", "error")
        return redirect(url_for("login"))
    except Exception:
        app.logger.exception("Google sign-in callback failed")
        flash("Google sign-in could not be completed. Please try again.", "error")
        return redirect(url_for("login"))

    email = _normalize_email(userinfo.get("email"))
    google_sub = (userinfo.get("sub") or "").strip()
    email_verified = userinfo.get("email_verified") in {True, "true", "True", 1, "1"}
    if not email or not google_sub or not email_verified:
        flash("Google did not provide a verified email address.", "error")
        return redirect(url_for("login"))

    allowed_domain = os.environ.get("GOOGLE_ALLOWED_DOMAIN", "").strip().lower().lstrip("@")
    if allowed_domain and email.rsplit("@", 1)[-1] != allowed_domain:
        flash(f"Use your @{allowed_domain} Google account.", "error")
        return redirect(url_for("login"))

    db = get_db()
    user = db.execute("SELECT * FROM users WHERE google_sub = ?", (google_sub,)).fetchone()
    if user and _normalize_email(user["email"]) != email:
        flash("This Google account no longer matches its authorized CRM email. Ask a manager to update it.", "error")
        return redirect(url_for("login"))

    if not user:
        user = db.execute("SELECT * FROM users WHERE email = ? COLLATE NOCASE", (email,)).fetchone()
        if user and user["google_sub"] and user["google_sub"] != google_sub:
            flash("That CRM email is already linked to another Google account.", "error")
            return redirect(url_for("login"))

    # Recovery/migration path for an existing manager whose old username was not
    # an email. It only activates for an explicitly configured bootstrap address.
    if not user and email in _bootstrap_emails():
        user = db.execute(
            "SELECT * FROM users WHERE role='manager' AND (email IS NULL OR TRIM(email)='') ORDER BY id LIMIT 1"
        ).fetchone()
        if not user and db.execute("SELECT COUNT(*) c FROM users").fetchone()["c"] == 0:
            display_name = (userinfo.get("name") or email.split("@", 1)[0]).strip()
            cursor = db.execute(
                "INSERT INTO users (username, password_hash, name, role, email, google_sub, last_login_at, created_at) "
                "VALUES (?, 'google-only', ?, 'manager', ?, ?, ?, ?)",
                (email, display_name, email, google_sub, now_str(), now_str()),
            )
            user = db.execute("SELECT * FROM users WHERE id=?", (cursor.lastrowid,)).fetchone()

    if not user:
        flash("This Google account is not authorized for QuantumLoop CRM. Ask a manager to add its email.", "error")
        return redirect(url_for("login"))

    try:
        db.execute(
            "UPDATE users SET email=?, google_sub=?, last_login_at=? WHERE id=?",
            (email, google_sub, now_str(), user["id"]),
        )
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        flash("This Google account conflicts with another CRM user. Ask a manager to correct the email.", "error")
        return redirect(url_for("login"))

    next_url = _safe_next_url(session.get("post_login_next")) or url_for("dashboard")
    session.clear()
    session.permanent = True
    session["user_id"] = user["id"]
    session["role"] = user["role"]
    g.user = db.execute("SELECT * FROM users WHERE id=?", (user["id"],)).fetchone()
    log_activity("Login with Google", "auth", user["id"], user["name"])
    return redirect(next_url)


@app.route("/logout")
def logout():
    user = current_user()
    if user:
        log_activity("Logout", "auth", user["id"], user["name"])
    session.clear()
    return redirect(url_for("login"))


# ---------- Users (manager only) ----------

@app.route("/users")
@manager_required
def users_list():
    db = get_db()
    users = db.execute("SELECT * FROM users ORDER BY role, name").fetchall()
    return render_template("users.html", users=users)


@app.route("/users/add", methods=["POST"])
@manager_required
def add_user():
    db = get_db()
    name = request.form.get("name", "").strip()
    email = _normalize_email(request.form.get("email"))
    role = request.form.get("role", "telecaller")

    if not name or not _valid_email(email) or role not in ROLES:
        flash("Enter a name, valid Google email and role.", "error")
        return redirect(url_for("users_list"))

    existing = db.execute(
        "SELECT id FROM users WHERE email=? COLLATE NOCASE OR username=? COLLATE NOCASE",
        (email, email),
    ).fetchone()
    if existing:
        flash(f"Google email '{email}' is already assigned. Update the existing user instead.", "error")
        return redirect(url_for("users_list"))

    db.execute(
        "INSERT INTO users (username, password_hash, name, role, email, created_at) "
        "VALUES (?, 'google-only', ?, ?, ?, ?)",
        (email, name, role, email, now_str()),
    )
    db.commit()
    flash(f"{name} can now sign in with {email}.", "success")
    return redirect(url_for("users_list"))


@app.route("/users/<int:user_id>/google-email", methods=["POST"])
@manager_required
def update_user_google_email(user_id):
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if not user:
        flash("User not found.", "error")
        return redirect(url_for("users_list"))

    email = _normalize_email(request.form.get("email"))
    if not _valid_email(email):
        flash("Enter a valid Google account email.", "error")
        return redirect(url_for("users_list"))
    existing = db.execute(
        "SELECT id FROM users WHERE email=? COLLATE NOCASE AND id<>?",
        (email, user_id),
    ).fetchone()
    if existing:
        flash(f"Google email '{email}' is already assigned to another user.", "error")
        return redirect(url_for("users_list"))

    old_email = _normalize_email(user["email"])
    unlink_google = old_email != email
    db.execute(
        "UPDATE users SET email=?, google_sub=CASE WHEN ? THEN NULL ELSE google_sub END WHERE id=?",
        (email, 1 if unlink_google else 0, user_id),
    )
    db.commit()
    flash(f"Google login email updated for {user['name']}.", "success")
    return redirect(url_for("users_list"))


@app.route("/users/<int:user_id>/delete", methods=["POST"])
@manager_required
def delete_user(user_id):
    if user_id == session.get("user_id"):
        flash("You can't delete your own account while logged in.", "error")
        return redirect(url_for("users_list"))
    db = get_db()
    db.execute("UPDATE leads SET owner_id = NULL WHERE owner_id = ?", (user_id,))
    db.execute("DELETE FROM users WHERE id = ?", (user_id,))
    db.commit()
    flash("User removed. Their leads are kept, now unassigned.", "success")
    return redirect(url_for("users_list"))


# ---------- Dashboard ----------

@app.route("/")
@login_required
def dashboard():
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)

    stage_counts = {s: 0 for s in STAGES}
    for row in db.execute(f"SELECT stage, COUNT(*) c FROM leads WHERE {clause} GROUP BY stage", params):
        if row["stage"] in stage_counts:
            stage_counts[row["stage"]] = row["c"]

    total_leads = db.execute(f"SELECT COUNT(*) c FROM leads WHERE {clause}", params).fetchone()["c"]

    today = today_str()
    week_ago = (datetime.now(APP_TIMEZONE).date() - timedelta(days=7)).strftime("%Y-%m-%d")

    added_today = db.execute(
        f"SELECT COUNT(*) c FROM leads WHERE {clause} AND date(created_at) = ?", [*params, today]
    ).fetchone()["c"]
    added_week = db.execute(
        f"SELECT COUNT(*) c FROM leads WHERE {clause} AND date(created_at) >= ?", [*params, week_ago]
    ).fetchone()["c"]

    won = stage_counts.get("Won", 0)
    conversion_rate = round((won / total_leads) * 100, 1) if total_leads else 0

    followups_today = db.execute(
        f"""SELECT leads.*, companies.name AS company_name FROM leads
            LEFT JOIN companies ON companies.id = leads.company_id
            WHERE {clause} AND follow_up_date = ? ORDER BY leads.name""", [*params, today]
    ).fetchall()
    followups_overdue = db.execute(
        f"""SELECT leads.*, companies.name AS company_name FROM leads
            LEFT JOIN companies ON companies.id = leads.company_id
            WHERE {clause} AND follow_up_date < ? AND follow_up_date IS NOT NULL AND follow_up_date != ''
            ORDER BY follow_up_date""",
        [*params, today],
    ).fetchall()

    team_stats = None
    monthly_stats = None
    field_visit_stats = None
    field_session_rows = None
    field_outcome_stats = None
    if user["role"] == "manager":
        team_stats = []
        telecallers = db.execute("SELECT * FROM users WHERE role = 'telecaller' ORDER BY name").fetchall()
        for tc in telecallers:
            calls_today = db.execute(
                "SELECT COUNT(*) c FROM call_logs WHERE user_id = ? AND date(created_at) = ?", (tc["id"], today)
            ).fetchone()["c"]
            leads_added = db.execute("SELECT COUNT(*) c FROM leads WHERE owner_id = ?", (tc["id"],)).fetchone()["c"]
            interested = db.execute(
                "SELECT COUNT(*) c FROM leads WHERE owner_id = ? AND stage = 'Interested'", (tc["id"],)
            ).fetchone()["c"]
            qualified = db.execute(
                "SELECT COUNT(*) c FROM leads WHERE owner_id = ? AND stage = 'Qualified'", (tc["id"],)
            ).fetchone()["c"]
            won_c = db.execute(
                "SELECT COUNT(*) c FROM leads WHERE owner_id = ? AND stage = 'Won'", (tc["id"],)
            ).fetchone()["c"]
            team_stats.append({
                "name": tc["name"], "calls_today": calls_today, "leads_added": leads_added,
                "interested": interested, "qualified": qualified, "won": won_c,
            })

        m_start = month_start_str()
        m_total = db.execute("SELECT COUNT(*) c FROM leads WHERE date(created_at) >= ?", (m_start,)).fetchone()["c"]
        m_interested = db.execute("SELECT COUNT(*) c FROM leads WHERE date(created_at) >= ? AND stage='Interested'", (m_start,)).fetchone()["c"]
        m_qualified = db.execute("SELECT COUNT(*) c FROM leads WHERE date(created_at) >= ? AND stage='Qualified'", (m_start,)).fetchone()["c"]
        m_won = db.execute("SELECT COUNT(*) c FROM leads WHERE date(created_at) >= ? AND stage='Won'", (m_start,)).fetchone()["c"]
        monthly_stats = {
            "total_leads": m_total,
            "interested_pct": round(m_interested / m_total * 100, 1) if m_total else 0,
            "qualified_pct": round(m_qualified / m_total * 100, 1) if m_total else 0,
            "won_pct": round(m_won / m_total * 100, 1) if m_total else 0,
            "conversion_rate": round(m_won / m_total * 100, 1) if m_total else 0,
        }

        field_visit_stats = {
            "scheduled_today": db.execute(
                "SELECT COUNT(*) c FROM field_visits WHERE scheduled_date=?", (today,)
            ).fetchone()["c"],
            "completed_today": db.execute(
                "SELECT COUNT(*) c FROM field_visits WHERE date(completed_at)=?", (today,)
            ).fetchone()["c"],
            "review_pending": db.execute(
                "SELECT COUNT(*) c FROM field_visits WHERE review_status='pending'"
            ).fetchone()["c"],
            "followups_due": db.execute(
                "SELECT COUNT(*) c FROM field_visits WHERE follow_up_date<=? AND follow_up_date!='' AND review_status='accepted'",
                (today,),
            ).fetchone()["c"],
        }
        field_session_rows = db.execute(
            """SELECT s.*, users.name AS executive_name,
                      COUNT(v.id) AS visit_count
               FROM field_work_sessions s
               LEFT JOIN users ON users.id=s.user_id
               LEFT JOIN field_visits v ON v.session_id=s.id
               WHERE s.work_date=?
               GROUP BY s.id ORDER BY s.started_at DESC""",
            (today,),
        ).fetchall()
        field_outcome_stats = db.execute(
            """SELECT outcome,COUNT(*) AS count FROM field_visits
               WHERE outcome IS NOT NULL AND outcome!='' AND date(completed_at)>=?
               GROUP BY outcome ORDER BY count DESC,outcome""",
            (m_start,),
        ).fetchall()

    return render_template(
        "dashboard.html",
        stage_counts=stage_counts,
        stages=STAGES,
        total_leads=total_leads,
        added_today=added_today,
        added_week=added_week,
        conversion_rate=conversion_rate,
        followups_today=followups_today,
        followups_overdue=followups_overdue,
        team_stats=team_stats,
        monthly_stats=monthly_stats,
        field_visit_stats=field_visit_stats,
        field_session_rows=field_session_rows,
        field_outcome_stats=field_outcome_stats,
    )


# ---------- Lead list ----------

@app.route("/leads")
@login_required
def leads_list():
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)

    stage_filter = request.args.get("stage", "")
    search = request.args.get("q", "").strip()
    tag_filter = request.args.get("tag", "").strip()
    interest_filter = request.args.get("interest", "").strip()

    query = f"""
        SELECT leads.*, users.name AS owner_name, companies.name AS company_name
        FROM leads
        LEFT JOIN users ON leads.owner_id = users.id
        LEFT JOIN companies ON leads.company_id = companies.id
        WHERE {clause}
    """
    if stage_filter:
        query += " AND stage = ?"
        params.append(stage_filter)
    if interest_filter:
        query += " AND interest_level = ?"
        params.append(interest_filter)
    if search:
        query += " AND (leads.name LIKE ? OR leads.phone LIKE ? OR leads.email LIKE ? OR companies.name LIKE ?)"
        like = f"%{search}%"
        params.extend([like, like, like, like])
    if tag_filter:
        query += " AND tags LIKE ?"
        params.append(f"%{tag_filter}%")
    query += " ORDER BY leads.created_at DESC"

    leads = db.execute(query, params).fetchall()

    return render_template(
        "leads_list.html",
        leads=leads,
        stages=STAGES,
        stage_filter=stage_filter,
        search=search,
        tag_filter=tag_filter,
        interest_filter=interest_filter,
        interest_levels=INTEREST_LEVELS,
        user=user,
    )


# ---------- Kanban ----------

@app.route("/leads/kanban")
@login_required
def leads_kanban():
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)

    query = f"""
        SELECT leads.*, users.name AS owner_name, companies.name AS company_name
        FROM leads
        LEFT JOIN users ON leads.owner_id = users.id
        LEFT JOIN companies ON leads.company_id = companies.id
        WHERE {clause}
        ORDER BY leads.created_at DESC
    """
    leads = db.execute(query, params).fetchall()

    columns = {s: [] for s in STAGES}
    for lead in leads:
        columns.setdefault(lead["stage"], []).append(lead)

    return render_template("kanban.html", columns=columns, stages=STAGES, user=user, time_ago=time_ago)


@app.route("/leads/<int:lead_id>/stage", methods=["POST"])
@login_required
def set_lead_stage(lead_id):
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)
    lead = db.execute(f"SELECT * FROM leads WHERE id = ? AND {clause}", [lead_id, *params]).fetchone()
    if lead is None:
        return jsonify({"ok": False, "error": "Lead not found or not visible to you."}), 404

    stage = request.json.get("stage") if request.is_json else request.form.get("stage")
    if stage not in STAGES:
        return jsonify({"ok": False, "error": "Invalid stage."}), 400
    if stage == "Lost":
        return jsonify({"ok": False, "error": "Open this lead to set a Lost reason before moving it to Lost."}), 400

    old_stage = lead["stage"]
    db.execute("UPDATE leads SET stage = ?, updated_at = ? WHERE id = ?", (stage, now_str(), lead_id))
    db.commit()
    log_lead_event(lead_id, "Stage Changed", f"{old_stage} → {stage}")
    log_activity("Stage Changed", "lead", lead_id, lead["name"], f"{old_stage} → {stage}")
    return jsonify({"ok": True})


# ---------- Company duplicate-check (Add Lead workflow, step 1) ----------

@app.route("/api/companies/check")
@login_required
def api_companies_check():
    db = get_db()
    name = request.args.get("name", "").strip()
    if not name or len(name) < 2:
        return jsonify({"matches": []})
    like = f"%{name}%"
    rows = db.execute(
        "SELECT id, name, industry, city FROM companies WHERE name LIKE ? ORDER BY name LIMIT 8", (like,)
    ).fetchall()
    return jsonify({"matches": [dict(r) for r in rows]})


@app.route("/api/companies/<int:company_id>/contacts")
@login_required
def api_company_contacts(company_id):
    db = get_db()
    rows = db.execute(
        "SELECT id, name, designation, phone, email, is_primary FROM contacts WHERE company_id = ? ORDER BY is_primary DESC, name",
        (company_id,),
    ).fetchall()
    return jsonify({"contacts": [dict(r) for r in rows]})


# ---------- Add lead (new company-centric workflow) ----------

@app.route("/leads/add", methods=["GET", "POST"])
@login_required
def add_lead():
    db = get_db()
    if request.method == "POST":
        user = current_user()
        f = request.form
        next_target = f.get("next_target", "").strip()

        company_id = f.get("company_id", "").strip()
        company_name = f.get("company_name", "").strip()
        contact_id = f.get("contact_id", "").strip()
        contact_name = f.get("contact_name", "").strip()

        if not company_id and not company_name:
            flash("Company name is required.", "error")
            return redirect(url_for("add_lead"))
        if not contact_id and not contact_name:
            flash("Contact name is required.", "error")
            return redirect(url_for("add_lead"))

        # Step 1/2 of the Add Lead workflow: reuse an existing company, or create one.
        gbp_link = f.get("gbp_link", "").strip()
        if company_id:
            company = db.execute("SELECT * FROM companies WHERE id = ?", (company_id,)).fetchone()
            if company is None:
                flash("Selected company not found.", "error")
                return redirect(url_for("add_lead"))
            company_id = company["id"]
            if gbp_link and not company["gbp_link"]:
                db.execute("UPDATE companies SET gbp_link = ? WHERE id = ?", (gbp_link, company_id))
        else:
            company = db.execute(
                "SELECT * FROM companies WHERE LOWER(TRIM(name))=LOWER(TRIM(?)) LIMIT 1", (company_name,)
            ).fetchone()
            if company:
                company_id = company["id"]
                company_name = company["name"]
                if gbp_link and not company["gbp_link"]:
                    db.execute("UPDATE companies SET gbp_link=? WHERE id=?", (gbp_link, company_id))
            else:
                cur = db.execute(
                    "INSERT INTO companies (name, industry, gbp_link, status, account_owner_id, created_at, updated_at) VALUES (?,?,?,?,?,?,?)",
                    (company_name, f.get("industry", "").strip(), gbp_link, "Prospect", user["id"], now_str(), now_str()),
                )
                company_id = cur.lastrowid
                log_activity("Company Created", "company", company_id, company_name)

        # Contact: reuse or create
        contact_phone = f.get("phone", "").strip()
        contact_alt_phone = f.get("alternate_phone", "").strip()
        contact_email = f.get("email", "").strip()
        designation = f.get("designation", "").strip()

        if contact_id:
            contact = db.execute("SELECT * FROM contacts WHERE id = ? AND company_id = ?", (contact_id, company_id)).fetchone()
            if contact is None:
                flash("Selected contact not found for this company.", "error")
                return redirect(url_for("add_lead"))
            contact_id = contact["id"]
            contact_name = contact["name"]
            contact_phone = contact_phone or contact["phone"]
            contact_email = contact_email or contact["email"]
            designation = designation or contact["designation"]
        else:
            contact = db.execute(
                """SELECT * FROM contacts WHERE company_id=? AND
                   (LOWER(TRIM(name))=LOWER(TRIM(?)) OR (?!='' AND phone=?) OR (?!='' AND LOWER(email)=LOWER(?)))
                   ORDER BY is_primary DESC LIMIT 1""",
                (company_id, contact_name, contact_phone, contact_phone, contact_email, contact_email),
            ).fetchone()
            if contact:
                contact_id = contact["id"]
                contact_name = contact["name"]
                contact_phone = contact_phone or contact["phone"]
                contact_email = contact_email or contact["email"]
                designation = designation or contact["designation"]
                db.execute(
                    "UPDATE contacts SET designation=?,phone=?,alternate_phone=?,email=?,updated_at=? WHERE id=?",
                    (designation, contact_phone, contact_alt_phone or contact["alternate_phone"], contact_email, now_str(), contact_id),
                )
            else:
                cur = db.execute(
                    "INSERT INTO contacts (company_id, name, designation, phone, alternate_phone, email, is_primary, created_at) VALUES (?,?,?,?,?,?,?,?)",
                    (company_id, contact_name, designation, contact_phone, contact_alt_phone, contact_email, 0, now_str()),
                )
                contact_id = cur.lastrowid
                log_activity("Contact Created", "contact", contact_id, contact_name)

        existing_lead = db.execute(
            """SELECT id,name FROM leads WHERE company_id=? AND
               (contact_id=? OR (?!='' AND phone=?)) ORDER BY id DESC LIMIT 1""",
            (company_id, contact_id, contact_phone, contact_phone),
        ).fetchone()
        if existing_lead:
            db.commit()
            flash(f"An existing lead for '{existing_lead['name']}' was reused to avoid a duplicate.", "success")
            if next_target == "field_visit" and user["role"] == "field_sales":
                return redirect(url_for("field_visit_new", lead_id=existing_lead["id"]))
            return redirect(url_for("lead_detail", lead_id=existing_lead["id"]))

        services = f.getlist("services_required")
        other_service = f.get("other_service_description", "").strip() if "Other" in services else ""

        cur = db.execute(
            """INSERT INTO leads (
                name, phone, email, source, stage, tags, follow_up_date, owner_id, created_at,
                company_id, contact_id, industry, designation, alternate_phone,
                website_status, website_url, website_quality, google_presence, social_presence,
                services_required, other_service_description, decision_maker, interest_level, priority,
                budget_range, estimated_project_value, expected_timeline,
                qualification_notes, requirement_notes, internal_notes, updated_at
            ) VALUES (?,?,?,?,?,?,?,?,?, ?,?,?,?,?, ?,?,?,?,?, ?,?,?,?,?, ?,?,?, ?,?,?,?)""",
            (
                contact_name, contact_phone, contact_email, f.get("source", "").strip(), "New",
                f.get("tags", "").strip(), f.get("follow_up_date", "").strip(), user["id"], now_str(),
                company_id, contact_id, f.get("industry", "").strip(), designation, contact_alt_phone,
                f.get("website_status", ""), f.get("website_url", "").strip(), f.get("website_quality", ""),
                f.get("google_presence", ""), f.get("social_presence", ""),
                ",".join(services), other_service, f.get("decision_maker", ""), f.get("interest_level", ""),
                f.get("priority", ""), f.get("budget_range", ""), f.get("estimated_project_value", "").strip(),
                f.get("expected_timeline", "").strip(), f.get("qualification_notes", "").strip(),
                f.get("requirement_notes", "").strip(), f.get("internal_notes", "").strip(), now_str(),
            ),
        )
        db.commit()
        lead_id = cur.lastrowid
        log_lead_event(lead_id, "Lead Created")
        log_activity("Lead Created", "lead", lead_id, contact_name, f"Company: {company_name or company_id}")
        flash(f"Lead for '{contact_name}' added.", "success")
        if next_target == "field_visit" and user["role"] == "field_sales":
            return redirect(url_for("field_visit_new", lead_id=lead_id))
        return redirect(url_for("lead_detail", lead_id=lead_id))

    companies = db.execute("SELECT id, name FROM companies ORDER BY name").fetchall()
    return render_template(
        "add_lead.html", stages=STAGES, companies=companies, services=SERVICES,
        website_status_opts=WEBSITE_STATUS, website_quality_opts=WEBSITE_QUALITY,
        google_presence_opts=GOOGLE_PRESENCE, social_presence_opts=SOCIAL_PRESENCE,
        decision_maker_opts=DECISION_MAKER_OPTIONS, interest_levels=INTEREST_LEVELS,
        priorities=PRIORITIES, budget_ranges=BUDGET_RANGES,
        next_target=request.args.get("next", ""),
    )


# ---------- CSV import ----------

REQUIRED_IMPORT_COLS = {"name", "phone", "company_name"}
OPTIONAL_IMPORT_COLS = {
    "email", "source", "tags", "industry", "designation", "alternate_phone",
    "interest_level", "priority", "budget_range", "decision_maker",
    "follow_up_date", "expected_timeline", "estimated_project_value",
    "requirement_notes", "internal_notes", "qualification_notes",
}


@app.route("/leads/import", methods=["POST"])
@login_required
def import_leads():
    db = get_db()
    user = current_user()
    file = request.files.get("csv_file")

    if not file or file.filename == "":
        flash("Please choose a CSV file to import.", "error")
        return redirect(url_for("leads_list"))

    if not file.filename.lower().endswith(".csv"):
        flash("That file doesn't look like a CSV. Export as .csv and try again.", "error")
        return redirect(url_for("leads_list"))

    try:
        raw = file.stream.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        flash("Couldn't read that file — please save it as UTF-8 CSV and try again.", "error")
        return redirect(url_for("leads_list"))

    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        flash("That CSV looks empty.", "error")
        return redirect(url_for("leads_list"))

    headers = {(h or "").strip().lower(): h for h in reader.fieldnames}
    missing = REQUIRED_IMPORT_COLS - set(headers.keys())
    if missing:
        flash(
            f"CSV is missing required column(s): {', '.join(sorted(missing))}. "
            f"Download the template below to see the expected format.",
            "error",
        )
        return redirect(url_for("leads_list"))

    def normalize_choice(value, valid_options):
        """Case-insensitive match against a controlled option list; blank if no match."""
        if not value:
            return ""
        for opt in valid_options:
            if opt.lower() == value.lower():
                return opt
        return ""

    imported = 0
    skipped_rows = []
    for i, row in enumerate(reader, start=2):  # row 1 is the header
        def get(col):
            key = headers.get(col)
            return (row.get(key) or "").strip() if key else ""

        name = get("name")
        phone = get("phone")
        company_name = get("company_name")
        if not name or not phone or not company_name:
            skipped_rows.append(str(i))
            continue

        # Reuse an existing company by name (case-insensitive), or create one — same as Add Lead.
        company = db.execute(
            "SELECT * FROM companies WHERE LOWER(name) = LOWER(?)", (company_name,)
        ).fetchone()
        industry = get("industry")
        if company is None:
            cur = db.execute(
                "INSERT INTO companies (name, industry, status, account_owner_id, created_at, updated_at) VALUES (?,?,?,?,?,?)",
                (company_name, industry, "Prospect", user["id"], now_str(), now_str()),
            )
            company_id = cur.lastrowid
            log_activity("Company Created", "company", company_id, company_name, "Imported via CSV")
        else:
            company_id = company["id"]

        designation = get("designation")
        alt_phone = get("alternate_phone")
        email = get("email")

        cur = db.execute(
            "INSERT INTO contacts (company_id, name, designation, phone, alternate_phone, email, is_primary, created_at) VALUES (?,?,?,?,?,?,?,?)",
            (company_id, name, designation, phone, alt_phone, email, 0, now_str()),
        )
        contact_id = cur.lastrowid

        cur = db.execute(
            """INSERT INTO leads (
                name, phone, email, source, stage, tags, follow_up_date, owner_id, created_at,
                company_id, contact_id, industry, designation, alternate_phone,
                decision_maker, interest_level, priority, budget_range,
                estimated_project_value, expected_timeline,
                qualification_notes, requirement_notes, internal_notes, updated_at
            ) VALUES (?,?,?,?,?,?,?,?,?, ?,?,?,?,?, ?,?,?,?, ?,?, ?,?,?,?)""",
            (
                name, phone, email, get("source"), "New", get("tags"), get("follow_up_date"),
                user["id"], now_str(),
                company_id, contact_id, industry, designation, alt_phone,
                normalize_choice(get("decision_maker"), DECISION_MAKER_OPTIONS),
                normalize_choice(get("interest_level"), INTEREST_LEVELS),
                normalize_choice(get("priority"), PRIORITIES),
                normalize_choice(get("budget_range"), BUDGET_RANGES),
                get("estimated_project_value"), get("expected_timeline"),
                get("qualification_notes"), get("requirement_notes"), get("internal_notes"), now_str(),
            ),
        )
        log_lead_event(cur.lastrowid, "Lead Created", "Imported via CSV")
        imported += 1

    db.commit()
    if imported:
        log_activity("Lead Created", "lead", None, "", f"Imported {imported} lead(s) via CSV")
        msg = f"Imported {imported} lead(s)."
        if skipped_rows:
            msg += f" Skipped {len(skipped_rows)} row(s) missing name/phone/company_name (row {', '.join(skipped_rows)})."
        flash(msg, "success")
    else:
        flash(f"No leads imported — every row was missing a name, phone, or company_name (rows {', '.join(skipped_rows)}).", "error")

    return redirect(url_for("leads_list"))


@app.route("/leads/template")
@login_required
def download_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads Template"

    headers = [
        "name", "phone", "company_name", "email", "source", "tags",
        "industry", "designation", "alternate_phone",
        "interest_level", "priority", "budget_range", "decision_maker",
        "follow_up_date", "expected_timeline", "estimated_project_value",
        "requirement_notes", "internal_notes", "qualification_notes",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="252D52", end_color="252D52", fill_type="solid")

    example = [
        "Ravi Kumar", "9876543210", "Acme Traders", "ravi@example.com", "Facebook", "hot, budget",
        "Retail", "Owner", "9876500000",
        "Hot", "High", "50k-1L", "Yes",
        "2026-08-20", "1-2 months", "75000",
        "Needs an online store", "Called twice, interested", "Budget confirmed",
    ]
    for col, value in enumerate(example, start=1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.font = Font(name="Arial", italic=True, color="9497B5")

    ws.cell(row=4, column=1, value="name, phone, and company_name are required for every row. All other columns are optional.").font = Font(name="Arial", size=10, italic=True)
    ws.cell(row=5, column=1, value=f"interest_level must be one of: {', '.join(INTEREST_LEVELS)}").font = Font(name="Arial", size=10, italic=True)
    ws.cell(row=6, column=1, value=f"priority must be one of: {', '.join(PRIORITIES)}").font = Font(name="Arial", size=10, italic=True)
    ws.cell(row=7, column=1, value=f"budget_range must be one of: {', '.join(BUDGET_RANGES)}").font = Font(name="Arial", size=10, italic=True)
    ws.cell(row=8, column=1, value="A company_name that already exists will be reused; otherwise a new company is created.").font = Font(name="Arial", size=10, italic=True)
    ws.cell(row=9, column=1, value="Delete the example row (row 2) before importing your real leads.").font = Font(name="Arial", size=10, italic=True)

    widths = [18, 14, 20, 24, 14, 18, 14, 16, 16, 12, 10, 12, 14, 14, 16, 20, 22, 22, 22]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment;filename=leads_import_template.xlsx"},
    )


# ---------- Export CSV ----------

@app.route("/leads/export")
@login_required
def export_leads():
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)
    leads = db.execute(
        f"""SELECT leads.*, companies.name AS company_name FROM leads
            LEFT JOIN companies ON companies.id = leads.company_id
            WHERE {clause} ORDER BY leads.created_at DESC""", params
    ).fetchall()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["company", "contact_name", "phone", "email", "source", "stage", "interest_level",
                      "priority", "budget_range", "tags", "follow_up_date", "created_at"])
    for l in leads:
        writer.writerow([l["company_name"], l["name"], l["phone"], l["email"], l["source"], l["stage"],
                          l["interest_level"], l["priority"], l["budget_range"], l["tags"],
                          l["follow_up_date"], l["created_at"]])

    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=leads_export.csv"},
    )


# ---------- Lead detail ----------

@app.route("/leads/<int:lead_id>")
@login_required
def lead_detail(lead_id):
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)
    lead = db.execute(
        f"""
        SELECT leads.*, users.name AS owner_name, companies.name AS company_name, companies.id AS company_pk
        FROM leads
        LEFT JOIN users ON leads.owner_id = users.id
        LEFT JOIN companies ON leads.company_id = companies.id
        WHERE leads.id = ? AND {clause}
        """,
        [lead_id, *params],
    ).fetchone()
    if lead is None:
        flash("Lead not found, or it's not visible to your account.", "error")
        return redirect(url_for("leads_list"))

    contact = None
    if lead["contact_id"]:
        contact = db.execute("SELECT * FROM contacts WHERE id = ?", (lead["contact_id"],)).fetchone()

    company = None
    company_contacts = []
    if lead["company_id"]:
        company = db.execute("SELECT * FROM companies WHERE id = ?", (lead["company_id"],)).fetchone()
        company_contacts = db.execute(
            "SELECT * FROM contacts WHERE company_id = ? ORDER BY is_primary DESC, name", (lead["company_id"],)
        ).fetchall()

    call_logs = db.execute(
        "SELECT * FROM call_logs WHERE lead_id = ? ORDER BY created_at DESC", (lead_id,)
    ).fetchall()

    timeline_events = db.execute(
        "SELECT * FROM lead_activities WHERE lead_id = ? ORDER BY created_at DESC, id DESC", (lead_id,)
    ).fetchall()

    services_selected = (lead["services_required"] or "").split(",") if lead["services_required"] else []
    telecallers = db.execute("SELECT id, name FROM users WHERE role = 'telecaller' ORDER BY name").fetchall()

    is_owner = lead["owner_id"] == user["id"]
    shared_with = db.execute(
        """SELECT lead_shares.user_id, users.name, users.role
           FROM lead_shares JOIN users ON lead_shares.user_id = users.id
           WHERE lead_shares.lead_id = ? ORDER BY users.name""",
        (lead_id,),
    ).fetchall()
    shareable_users = []
    if is_owner:
        already_shared_ids = {row["user_id"] for row in shared_with}
        if user["role"] == "manager":
            # Managers already see every telecaller's leads automatically, so sharing
            # to a telecaller here is only useful for handing a manager's own lead to
            # a telecaller to work; sharing to another manager is also allowed.
            all_users = db.execute("SELECT id, name, role FROM users ORDER BY name").fetchall()
        else:
            # Telecallers can only send leads to other telecallers. Managers already
            # see every telecaller lead, so a telecaller "sharing" to a manager would
            # be meaningless — leave managers out of the list entirely.
            all_users = db.execute("SELECT id, name, role FROM users WHERE role = 'telecaller' ORDER BY name").fetchall()
        shareable_users = [u for u in all_users if u["id"] != user["id"] and u["id"] not in already_shared_ids]

    return render_template(
        "lead_detail.html", lead=lead, contact=contact, company=company, company_contacts=company_contacts,
        call_logs=call_logs, stages=STAGES,
        timeline_events=timeline_events, services=SERVICES, services_selected=services_selected,
        website_status_opts=WEBSITE_STATUS, website_quality_opts=WEBSITE_QUALITY,
        google_presence_opts=GOOGLE_PRESENCE, social_presence_opts=SOCIAL_PRESENCE,
        decision_maker_opts=DECISION_MAKER_OPTIONS, interest_levels=INTEREST_LEVELS,
        priorities=PRIORITIES, budget_ranges=BUDGET_RANGES, lost_reasons=LOST_REASONS,
        telecallers=telecallers, is_owner=is_owner, shared_with=shared_with, shareable_users=shareable_users,
    )


@app.route("/leads/<int:lead_id>/link-company", methods=["POST"])
@login_required
def link_company_to_lead(lead_id):
    db = get_db()
    user = current_user()
    lead = db.execute("SELECT * FROM leads WHERE id = ?", (lead_id,)).fetchone()
    if lead is None:
        flash("Lead not found.", "error")
        return redirect(url_for("leads_list"))

    f = request.form
    company_id = f.get("company_id", "").strip()
    company_name = f.get("company_name", "").strip()

    if not company_id and not company_name:
        flash("Enter a company name to link.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))

    if company_id:
        company = db.execute("SELECT * FROM companies WHERE id = ?", (company_id,)).fetchone()
        if company is None:
            flash("Selected company not found.", "error")
            return redirect(url_for("lead_detail", lead_id=lead_id))
        company_id = company["id"]
    else:
        cur = db.execute(
            "INSERT INTO companies (name, status, account_owner_id, created_at, updated_at) VALUES (?,?,?,?,?)",
            (company_name, "Prospect", user["id"], now_str(), now_str()),
        )
        company_id = cur.lastrowid
        log_activity("Company Created", "company", company_id, company_name)

    # This lead already has a name/phone/email of its own — turn that into (or reuse) a
    # contact record under the newly linked company, so the contacts list is populated too.
    contact_id = lead["contact_id"]
    if not contact_id:
        cur = db.execute(
            "INSERT INTO contacts (company_id, name, designation, phone, alternate_phone, email, is_primary, created_at) VALUES (?,?,?,?,?,?,?,?)",
            (company_id, lead["name"], lead["designation"], lead["phone"], lead["alternate_phone"], lead["email"], 1, now_str()),
        )
        contact_id = cur.lastrowid
        log_activity("Contact Created", "contact", contact_id, lead["name"])

    db.execute("UPDATE leads SET company_id = ?, contact_id = ?, updated_at = ? WHERE id = ?", (company_id, contact_id, now_str(), lead_id))
    db.commit()
    log_lead_event(lead_id, "Linked to Company", company_name or "")
    flash("Lead linked to company.", "success")
    return redirect(url_for("lead_detail", lead_id=lead_id))


@app.route("/leads/<int:lead_id>/update", methods=["POST"])
@login_required
def update_lead(lead_id):
    db = get_db()
    user = current_user()
    f = request.form
    clause, params = visible_leads_clause(user)
    lead = db.execute(f"SELECT * FROM leads WHERE id = ? AND {clause}", [lead_id, *params]).fetchone()
    if lead is None:
        flash("Lead not found, or it's not visible to your account.", "error")
        return redirect(url_for("leads_list"))

    stage = f.get("stage", lead["stage"])
    lost_reason = f.get("lost_reason", "").strip()
    if stage == "Lost" and not lost_reason:
        flash("Please choose a Lost reason before setting this lead to Lost.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))
    if stage != "Lost":
        lost_reason = ""

    services = f.getlist("services_required")
    other_service = f.get("other_service_description", "").strip() if "Other" in services else ""

    db.execute(
        """UPDATE leads SET
            stage=?, tags=?, follow_up_date=?, industry=?, designation=?, alternate_phone=?,
            website_status=?, website_url=?, website_quality=?, google_presence=?, social_presence=?,
            services_required=?, other_service_description=?, decision_maker=?, interest_level=?, priority=?,
            budget_range=?, estimated_project_value=?, expected_timeline=?, lost_reason=?,
            qualification_notes=?, requirement_notes=?, internal_notes=?, updated_at=?
        WHERE id=?""",
        (
            stage, f.get("tags", "").strip(), f.get("follow_up_date", "").strip(),
            f.get("industry", "").strip(), f.get("designation", "").strip(), f.get("alternate_phone", "").strip(),
            f.get("website_status", ""), f.get("website_url", "").strip(), f.get("website_quality", ""),
            f.get("google_presence", ""), f.get("social_presence", ""),
            ",".join(services), other_service, f.get("decision_maker", ""), f.get("interest_level", ""),
            f.get("priority", ""), f.get("budget_range", ""), f.get("estimated_project_value", "").strip(),
            f.get("expected_timeline", "").strip(), lost_reason,
            f.get("qualification_notes", "").strip(), f.get("requirement_notes", "").strip(),
            f.get("internal_notes", "").strip(), now_str(), lead_id,
        ),
    )
    db.commit()

    if stage != lead["stage"]:
        detail = f"{lead['stage']} → {stage}" + (f" ({lost_reason})" if stage == "Lost" else "")
        log_lead_event(lead_id, "Stage Changed", detail)
        log_activity("Stage Changed", "lead", lead_id, lead["name"], detail)
    else:
        log_lead_event(lead_id, "Lead Edited")
        log_activity("Lead Edited", "lead", lead_id, lead["name"])

    flash("Lead updated.", "success")
    return redirect(url_for("lead_detail", lead_id=lead_id))


@app.route("/leads/<int:lead_id>/call", methods=["POST"])
@login_required
def add_call_log(lead_id):
    db = get_db()
    user = current_user()
    clause, params = visible_leads_clause(user)
    lead_ok = db.execute(f"SELECT 1 FROM leads WHERE id = ? AND {clause}", [lead_id, *params]).fetchone()
    if lead_ok is None:
        flash("Lead not found, or it's not visible to your account.", "error")
        return redirect(url_for("leads_list"))
    outcome = request.form.get("outcome", "").strip()
    note = request.form.get("note", "").strip()

    db.execute(
        "INSERT INTO call_logs (lead_id, outcome, note, user_id, created_at) VALUES (?, ?, ?, ?, ?)",
        (lead_id, outcome, note, user["id"], now_str()),
    )
    db.commit()
    log_lead_event(lead_id, "Called", f"{outcome}" + (f" — {note}" if note else ""))
    lead = db.execute("SELECT name FROM leads WHERE id=?", (lead_id,)).fetchone()
    log_activity("Lead Edited", "lead", lead_id, lead["name"] if lead else "", f"Call logged: {outcome}")
    flash("Call logged.", "success")
    return redirect(url_for("lead_detail", lead_id=lead_id))


@app.route("/leads/<int:lead_id>/delete", methods=["POST"])
@login_required
def delete_lead(lead_id):
    db = get_db()
    user = current_user()
    lead = db.execute("SELECT * FROM leads WHERE id = ?", (lead_id,)).fetchone()
    if lead is None:
        flash("Lead not found.", "error")
        return redirect(url_for("leads_list"))
    if lead["owner_id"] != user["id"] and user["role"] != "manager":
        flash("Only the owner or a manager can delete this lead.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))
    db.execute("DELETE FROM lead_shares WHERE lead_id = ?", (lead_id,))
    db.execute("DELETE FROM call_logs WHERE lead_id = ?", (lead_id,))
    db.execute("DELETE FROM lead_activities WHERE lead_id = ?", (lead_id,))
    db.execute("DELETE FROM leads WHERE id = ?", (lead_id,))
    db.commit()
    log_activity("Lead Deleted", "lead", lead_id, lead["name"] if lead else "")
    flash("Lead deleted.", "success")
    return redirect(url_for("leads_list"))


@app.route("/leads/<int:lead_id>/assign", methods=["POST"])
@manager_required
def assign_lead(lead_id):
    """Sets initial ownership for a lead that doesn't have an owner yet (e.g. a
    freshly imported/unassigned lead). Once a lead has an owner, ownership stays
    put — use /leads/<id>/share to grant another person view+edit access instead."""
    db = get_db()
    lead = db.execute("SELECT name, owner_id FROM leads WHERE id = ?", (lead_id,)).fetchone()
    if lead is None:
        flash("Lead not found.", "error")
        return redirect(url_for("leads_list"))
    if lead["owner_id"]:
        flash("This lead already has an owner. Use \u201cShare access\u201d to give someone else access instead of reassigning it.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))

    owner_id = request.form.get("owner_id") or None
    db.execute("UPDATE leads SET owner_id = ? WHERE id = ?", (owner_id, lead_id))
    db.commit()
    owner = db.execute("SELECT name FROM users WHERE id = ?", (owner_id,)).fetchone() if owner_id else None
    log_lead_event(lead_id, "Lead Assigned", owner["name"] if owner else "Unassigned")
    log_activity("Lead Assigned", "lead", lead_id, lead["name"] if lead else "", owner["name"] if owner else "Unassigned")
    flash("Lead assignment updated.", "success")
    return redirect(url_for("lead_detail", lead_id=lead_id))


@app.route("/leads/<int:lead_id>/share", methods=["POST"])
@login_required
def share_lead(lead_id):
    """Grants another user (telecaller or manager) view+edit access to this lead,
    without changing who owns it and without creating a copy. Only the lead's
    owner can grant access to it."""
    db = get_db()
    user = current_user()
    lead = db.execute("SELECT * FROM leads WHERE id = ?", (lead_id,)).fetchone()
    if lead is None:
        flash("Lead not found.", "error")
        return redirect(url_for("leads_list"))
    if lead["owner_id"] != user["id"]:
        flash("Only the lead's owner can share access to it.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))

    grantee_id = request.form.get("user_id")
    if not grantee_id:
        flash("Choose a person to share this lead with.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))
    grantee_id = int(grantee_id)
    if grantee_id == user["id"]:
        flash("You already have access to your own lead.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))

    grantee = db.execute("SELECT id, name, role FROM users WHERE id = ?", (grantee_id,)).fetchone()
    if grantee is None:
        flash("That user doesn't exist.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))
    if user["role"] == "telecaller" and grantee["role"] == "manager":
        flash("Telecallers can only share leads with other telecallers — managers already see every telecaller's leads.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))

    db.execute(
        "INSERT OR IGNORE INTO lead_shares (lead_id, user_id, granted_by, created_at) VALUES (?, ?, ?, ?)",
        (lead_id, grantee_id, user["id"], now_str()),
    )
    db.commit()
    log_lead_event(lead_id, "Access Granted", grantee["name"])
    log_activity("Access Granted", "lead", lead_id, lead["name"], f"Shared with {grantee['name']}")
    flash(f"{grantee['name']} can now see and edit this lead.", "success")
    return redirect(url_for("lead_detail", lead_id=lead_id))


@app.route("/leads/<int:lead_id>/unshare/<int:user_id>", methods=["POST"])
@login_required
def unshare_lead(lead_id, user_id):
    """Revokes a previously granted access. Only the lead's owner can revoke it."""
    db = get_db()
    user = current_user()
    lead = db.execute("SELECT * FROM leads WHERE id = ?", (lead_id,)).fetchone()
    if lead is None:
        flash("Lead not found.", "error")
        return redirect(url_for("leads_list"))
    if lead["owner_id"] != user["id"]:
        flash("Only the lead's owner can revoke access.", "error")
        return redirect(url_for("lead_detail", lead_id=lead_id))

    revoked = db.execute("SELECT name FROM users WHERE id = ?", (user_id,)).fetchone()
    db.execute("DELETE FROM lead_shares WHERE lead_id = ? AND user_id = ?", (lead_id, user_id))
    db.commit()
    log_lead_event(lead_id, "Access Revoked", revoked["name"] if revoked else "")
    log_activity("Access Revoked", "lead", lead_id, lead["name"], f"Revoked from {revoked['name']}" if revoked else "")
    flash("Access revoked.", "success")
    return redirect(url_for("lead_detail", lead_id=lead_id))


# ---------- Companies (the center of everything) ----------

@app.route("/contacts")
@login_required
def contacts_list():
    db = get_db()
    search = request.args.get("q", "").strip()
    query = """
        SELECT companies.*, COUNT(DISTINCT contacts.id) AS contact_count
        FROM companies LEFT JOIN contacts ON contacts.company_id = companies.id
    """
    params = []
    if search:
        query += " WHERE companies.name LIKE ?"
        params.append(f"%{search}%")
    query += " GROUP BY companies.id ORDER BY companies.name"
    companies = db.execute(query, params).fetchall()
    users = db.execute("SELECT id, name FROM users ORDER BY name").fetchall()
    return render_template("contacts_list.html", companies=companies, search=search,
                            statuses=COMPANY_STATUSES, users=users)


@app.route("/contacts/companies/add", methods=["POST"])
@login_required
def add_company():
    db = get_db()
    user = current_user()
    name = request.form.get("name", "").strip()
    if not name:
        flash("Company name is required.", "error")
        return redirect(url_for("contacts_list"))

    f = request.form
    db.execute(
        """INSERT INTO companies (
            name, industry, phone, alternate_phone, email, website, gbp_link, notes, billing_address, address,
            city, state, pincode, country, gstin, pan, cin, status, account_owner_id, created_at, updated_at
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            name, f.get("industry", "").strip(), f.get("phone", "").strip(), f.get("alternate_phone", "").strip(),
            f.get("email", "").strip(), f.get("website", "").strip(), f.get("gbp_link", "").strip(), f.get("notes", "").strip(),
            f.get("billing_address", "").strip(), f.get("address", "").strip(), f.get("city", "").strip(),
            f.get("state", "").strip(), f.get("pincode", "").strip(), f.get("country", "India").strip() or "India",
            f.get("gstin", "").strip().upper(), f.get("pan", "").strip().upper(), f.get("cin", "").strip().upper(),
            f.get("status", "Prospect"), f.get("account_owner_id") or user["id"], now_str(), now_str(),
        ),
    )
    db.commit()
    company_id = db.execute("SELECT last_insert_rowid() id").fetchone()["id"]
    log_activity("Company Created", "company", company_id, name)
    flash(f"Company '{name}' added.", "success")
    return redirect(url_for("contacts_list"))


@app.route("/contacts/companies/<int:company_id>")
@login_required
def company_detail(company_id):
    db = get_db()
    company = db.execute("SELECT * FROM companies WHERE id = ?", (company_id,)).fetchone()
    if company is None:
        flash("Company not found.", "error")
        return redirect(url_for("contacts_list"))
    contacts = db.execute(
        "SELECT * FROM contacts WHERE company_id = ? ORDER BY is_primary DESC, name", (company_id,)
    ).fetchall()
    leads = db.execute(
        """SELECT leads.*, users.name AS owner_name FROM leads
           LEFT JOIN users ON users.id = leads.owner_id
           WHERE company_id = ? ORDER BY leads.created_at DESC""", (company_id,)
    ).fetchall()
    invoices = db.execute(
        "SELECT * FROM invoices WHERE company_id = ? ORDER BY id DESC", (company_id,)
    ).fetchall()
    activity = db.execute(
        "SELECT * FROM activity_logs WHERE entity_type = 'company' AND entity_id = ? ORDER BY created_at DESC LIMIT 50",
        (company_id,),
    ).fetchall()
    return render_template(
        "company_detail.html", company=company, contacts=contacts, leads=leads, invoices=invoices,
        activity=activity, statuses=COMPANY_STATUSES,
    )


@app.route("/contacts/companies/<int:company_id>/edit", methods=["POST"])
@login_required
def edit_company(company_id):
    db = get_db()
    company = db.execute("SELECT * FROM companies WHERE id = ?", (company_id,)).fetchone()
    if company is None:
        flash("Company not found.", "error")
        return redirect(url_for("contacts_list"))
    f = request.form
    name = f.get("name", "").strip()
    if not name:
        flash("Company name is required.", "error")
        return redirect(url_for("company_detail", company_id=company_id))
    db.execute(
        """UPDATE companies SET name=?, industry=?, phone=?, alternate_phone=?, email=?, website=?, gbp_link=?, notes=?,
            billing_address=?, address=?, city=?, state=?, pincode=?, country=?, gstin=?, pan=?, cin=?, status=?, updated_at=?
        WHERE id=?""",
        (
            name, f.get("industry", "").strip(), f.get("phone", "").strip(), f.get("alternate_phone", "").strip(),
            f.get("email", "").strip(), f.get("website", "").strip(), f.get("gbp_link", "").strip(), f.get("notes", "").strip(),
            f.get("billing_address", "").strip(), f.get("address", "").strip(), f.get("city", "").strip(),
            f.get("state", "").strip(), f.get("pincode", "").strip(), f.get("country", "India").strip() or "India",
            f.get("gstin", "").strip().upper(), f.get("pan", "").strip().upper(), f.get("cin", "").strip().upper(),
            f.get("status", "Prospect"), now_str(), company_id,
        ),
    )
    db.commit()
    log_activity("Company Updated", "company", company_id, name)
    flash("Company updated.", "success")
    return redirect(url_for("company_detail", company_id=company_id))


@app.route("/contacts/companies/<int:company_id>/delete", methods=["POST"])
@login_required
def delete_company(company_id):
    db = get_db()
    company = db.execute("SELECT name FROM companies WHERE id = ?", (company_id,)).fetchone()
    db.execute("DELETE FROM contacts WHERE company_id = ?", (company_id,))
    db.execute("UPDATE leads SET company_id = NULL WHERE company_id = ?", (company_id,))
    db.execute("DELETE FROM companies WHERE id = ?", (company_id,))
    db.commit()
    log_activity("Company Deleted", "company", company_id, company["name"] if company else "")
    flash("Company and its contacts deleted.", "success")
    return redirect(url_for("contacts_list"))


@app.route("/contacts/companies/<int:company_id>/contacts/add", methods=["POST"])
@login_required
def add_contact(company_id):
    db = get_db()
    name = request.form.get("name", "").strip()
    next_lead_id = request.form.get("next_lead_id", "").strip()
    redirect_target = url_for("lead_detail", lead_id=next_lead_id) if next_lead_id else url_for("company_detail", company_id=company_id)

    if not name:
        flash("Contact name is required.", "error")
        return redirect(redirect_target)

    f = request.form
    is_primary = 1 if f.get("is_primary") else 0
    if is_primary:
        db.execute("UPDATE contacts SET is_primary = 0 WHERE company_id = ?", (company_id,))

    db.execute(
        """INSERT INTO contacts (company_id, name, designation, phone, alternate_phone, email, linkedin, is_primary, notes, created_at)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (
            company_id, name, f.get("designation", "").strip(), f.get("phone", "").strip(),
            f.get("alternate_phone", "").strip(), f.get("email", "").strip(), f.get("linkedin", "").strip(),
            is_primary, f.get("notes", "").strip(), now_str(),
        ),
    )
    db.commit()
    contact_id = db.execute("SELECT last_insert_rowid() id").fetchone()["id"]
    log_activity("Contact Created", "contact", contact_id, name)
    flash(f"Contact '{name}' added.", "success")
    return redirect(redirect_target)


@app.route("/contacts/<int:contact_id>/delete", methods=["POST"])
@login_required
def delete_contact(contact_id):
    db = get_db()
    row = db.execute("SELECT company_id, name FROM contacts WHERE id = ?", (contact_id,)).fetchone()
    next_lead_id = request.form.get("next_lead_id", "").strip()
    db.execute("UPDATE leads SET contact_id = NULL WHERE contact_id = ?", (contact_id,))
    db.execute("DELETE FROM contacts WHERE id = ?", (contact_id,))
    db.commit()
    if row:
        log_activity("Contact Deleted", "contact", contact_id, row["name"])
    flash("Contact deleted.", "success")
    if next_lead_id:
        return redirect(url_for("lead_detail", lead_id=next_lead_id))
    return redirect(url_for("company_detail", company_id=row["company_id"] if row else 0))


# ---------- Global search ----------

@app.route("/search")
@login_required
def global_search():
    db = get_db()
    q = request.args.get("q", "").strip()
    results = {"companies": [], "contacts": [], "leads": [], "invoices": []}
    if q and len(q) >= 2:
        like = f"%{q}%"
        results["companies"] = db.execute(
            "SELECT * FROM companies WHERE name LIKE ? OR gstin LIKE ? OR email LIKE ? OR phone LIKE ? ORDER BY name LIMIT 25",
            (like, like, like, like),
        ).fetchall()
        results["contacts"] = db.execute(
            """SELECT contacts.*, companies.name AS company_name FROM contacts
               LEFT JOIN companies ON companies.id = contacts.company_id
               WHERE contacts.name LIKE ? OR contacts.phone LIKE ? OR contacts.email LIKE ? ORDER BY contacts.name LIMIT 25""",
            (like, like, like),
        ).fetchall()
        results["leads"] = db.execute(
            """SELECT leads.*, companies.name AS company_name FROM leads
               LEFT JOIN companies ON companies.id = leads.company_id
               WHERE leads.name LIKE ? OR leads.phone LIKE ? OR leads.email LIKE ? OR companies.name LIKE ?
               ORDER BY leads.created_at DESC LIMIT 25""",
            (like, like, like, like),
        ).fetchall()
        user = current_user()
        if user["role"] == "manager":
            results["invoices"] = db.execute(
                "SELECT * FROM invoices WHERE invoice_no LIKE ? OR bill_name LIKE ? ORDER BY id DESC LIMIT 25",
                (like, like),
            ).fetchall()
    return render_template("search_results.html", q=q, results=results)


# ---------- Activity Logs page (manager only) ----------

@app.route("/activity-logs")
@manager_required
def activity_logs_page():
    db = get_db()
    user_filter = request.args.get("user", "").strip()
    action_filter = request.args.get("action", "").strip()
    date_filter = request.args.get("date", "").strip()

    query = "SELECT * FROM activity_logs WHERE 1=1"
    params = []
    if user_filter:
        query += " AND user_id = ?"
        params.append(user_filter)
    if action_filter:
        query += " AND action = ?"
        params.append(action_filter)
    if date_filter:
        query += " AND date(created_at) = ?"
        params.append(date_filter)
    query += " ORDER BY created_at DESC LIMIT 300"

    logs = db.execute(query, params).fetchall()
    users = db.execute("SELECT id, name FROM users ORDER BY name").fetchall()
    actions = [r["action"] for r in db.execute("SELECT DISTINCT action FROM activity_logs ORDER BY action")]

    return render_template(
        "activity_logs.html", logs=logs, users=users, actions=actions,
        user_filter=user_filter, action_filter=action_filter, date_filter=date_filter,
    )


# ---------- Field Sales ----------

def _active_field_session(db, user_id):
    _reconcile_field_sessions(db)
    return db.execute(
        "SELECT * FROM field_work_sessions WHERE user_id=? AND work_date=? AND status='active'",
        (user_id, today_str()),
    ).fetchone()


def _session_duration(started_at, ended_at=None):
    if not started_at:
        return "—"
    start = datetime.strptime(started_at, "%Y-%m-%d %H:%M:%S")
    end = datetime.strptime(ended_at, "%Y-%m-%d %H:%M:%S") if ended_at else datetime.now(APP_TIMEZONE).replace(tzinfo=None)
    total_minutes = max(0, int((end - start).total_seconds() // 60))
    return f"{total_minutes // 60}h {total_minutes % 60}m"


def _get_field_visit(db, visit_id):
    return db.execute(
        """SELECT v.*, leads.name AS lead_name, leads.phone AS lead_phone,
                  leads.email AS lead_email, leads.designation AS lead_designation, leads.contact_id AS lead_contact_id,
                  companies.name AS company_name, users.name AS executive_name,
                  creators.name AS creator_name
           FROM field_visits v
           LEFT JOIN leads ON leads.id=v.lead_id
           LEFT JOIN companies ON companies.id=v.company_id
           LEFT JOIN users ON users.id=v.assigned_to
           LEFT JOIN users creators ON creators.id=v.created_by
           WHERE v.id=?""",
        (visit_id,),
    ).fetchone()


def _can_access_field_visit(user, visit):
    return bool(visit and (user["role"] == "manager" or visit["assigned_to"] == user["id"]))


def _log_field_visit_event(visit, event, detail=""):
    if visit["lead_id"]:
        log_lead_event(visit["lead_id"], event, detail)
    if visit["company_id"]:
        log_activity(event, "company", visit["company_id"], visit["company_name"] or "", detail)
    log_activity(event, "field_visit", visit["id"], visit["company_name"] or visit["lead_name"] or f"Visit {visit['id']}", detail)


def _photo_exists(db, visit_id, photo_type):
    return db.execute(
        "SELECT 1 FROM field_visit_photos WHERE visit_id=? AND photo_type=? LIMIT 1",
        (visit_id, photo_type),
    ).fetchone() is not None


def _field_photo_dir():
    return os.environ.get("FIELD_PHOTO_DIR", "").strip() or os.path.join(app.instance_path, "field_visit_photos")


def _save_camera_photo(db, visit_id, photo_type, data_url):
    if not data_url:
        return None
    prefix = "data:image/jpeg;base64,"
    if not data_url.startswith(prefix):
        raise ValueError("Camera evidence must be a JPEG captured by the visit form.")
    try:
        raw = base64.b64decode(data_url[len(prefix):], validate=True)
    except (ValueError, binascii.Error):
        raise ValueError("The captured camera image could not be read.")
    if len(raw) > 8 * 1024 * 1024:
        raise ValueError("The captured image is larger than 8 MB.")
    if len(raw) < 4 or not raw.startswith(b"\xff\xd8\xff"):
        raise ValueError("The captured evidence is not a valid JPEG image.")

    upload_dir = _field_photo_dir()
    os.makedirs(upload_dir, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}.jpg"
    with open(os.path.join(upload_dir, stored_name), "wb") as output:
        output.write(raw)
    db.execute(
        """INSERT INTO field_visit_photos
           (visit_id,photo_type,stored_name,mime_type,byte_size,sha256,captured_at,created_at)
           VALUES (?,?,?,?,?,?,?,?)""",
        (visit_id, photo_type, stored_name, "image/jpeg", len(raw), hashlib.sha256(raw).hexdigest(), now_str(), now_str()),
    )
    return stored_name


@app.route("/field-visits")
@field_sales_or_manager_required
def field_visits():
    db = get_db()
    user = current_user()
    session_filter = request.args.get("session", "").strip()
    query = """SELECT v.*, leads.name AS lead_name, companies.name AS company_name,
                      users.name AS executive_name
               FROM field_visits v
               LEFT JOIN leads ON leads.id=v.lead_id
               LEFT JOIN companies ON companies.id=v.company_id
               LEFT JOIN users ON users.id=v.assigned_to"""
    params = []
    if user["role"] == "field_sales":
        query += " WHERE v.assigned_to=?"
        params.append(user["id"])
    elif session_filter:
        query += " WHERE v.session_id=?"
        params.append(session_filter)
    query += " ORDER BY v.scheduled_date DESC, v.scheduled_time DESC, v.id DESC"
    visits = db.execute(query, params).fetchall()

    session_today = None
    sessions = []
    if user["role"] == "field_sales":
        session_today = db.execute(
            "SELECT * FROM field_work_sessions WHERE user_id=? AND work_date=?", (user["id"], today_str())
        ).fetchone()
    else:
        sessions = db.execute(
            """SELECT s.*, users.name AS executive_name, COUNT(v.id) AS visit_count
               FROM field_work_sessions s
               LEFT JOIN users ON users.id=s.user_id
               LEFT JOIN field_visits v ON v.session_id=s.id
               GROUP BY s.id ORDER BY s.work_date DESC, s.started_at DESC LIMIT 100"""
        ).fetchall()
    return render_template(
        "field_visits.html", visits=visits, session_today=session_today, sessions=sessions,
        session_duration=_session_duration, today=today_str(), session_filter=session_filter,
    )


@app.route("/field-work/start", methods=["POST"])
@field_sales_or_manager_required
def field_work_start():
    user = current_user()
    if user["role"] != "field_sales":
        flash("Only a Field Sales Executive can start field work.", "error")
        return redirect(url_for("field_visits"))
    db = get_db()
    existing = db.execute(
        "SELECT * FROM field_work_sessions WHERE user_id=? AND work_date=?", (user["id"], today_str())
    ).fetchone()
    if existing:
        flash("Today's field-work session has already been started.", "error")
        return redirect(url_for("field_visits"))
    if datetime.now(APP_TIMEZONE).time() >= FIELD_WORK_AUTO_END:
        flash("A field-work session cannot be started after the 6:30 p.m. automatic close time.", "error")
        return redirect(url_for("field_visits"))
    db.execute(
        "INSERT INTO field_work_sessions (user_id,work_date,started_at,status,created_at) VALUES (?,?,?,?,?)",
        (user["id"], today_str(), now_str(), "active", now_str()),
    )
    db.commit()
    log_activity("Field Work Started", "field_session", None, user["name"], today_str())
    flash("Field work started. Check in when you arrive at your first visit.", "success")
    return redirect(url_for("field_visits"))


@app.route("/field-work/end", methods=["POST"])
@field_sales_or_manager_required
def field_work_end():
    user = current_user()
    if user["role"] != "field_sales":
        flash("Only a Field Sales Executive can end field work.", "error")
        return redirect(url_for("field_visits"))
    db = get_db()
    work_session = _active_field_session(db, user["id"])
    if not work_session:
        flash("There is no active field-work session.", "error")
        return redirect(url_for("field_visits"))
    open_visit = db.execute(
        "SELECT id FROM field_visits WHERE session_id=? AND status='checked_in'", (work_session["id"],)
    ).fetchone()
    if open_visit:
        flash("Complete your checked-in visit before ending field work.", "error")
        return redirect(url_for("field_visit_detail", visit_id=open_visit["id"]))
    db.execute(
        "UPDATE field_work_sessions SET ended_at=?,end_source='executive',status='closed' WHERE id=?",
        (now_str(), work_session["id"]),
    )
    db.commit()
    log_activity("Field Work Ended", "field_session", work_session["id"], user["name"], today_str())
    flash("Field work ended for today.", "success")
    return redirect(url_for("field_visits"))


@app.route("/field-visits/new", methods=["GET", "POST"])
@field_sales_or_manager_required
def field_visit_new():
    db = get_db()
    user = current_user()
    if request.method == "POST":
        f = request.form
        if user["role"] == "manager":
            assigned_to = f.get("assigned_to", "").strip()
            executive = db.execute("SELECT * FROM users WHERE id=? AND role='field_sales'", (assigned_to,)).fetchone()
            if not executive:
                flash("Choose a valid Field Sales Executive.", "error")
                return redirect(url_for("field_visit_new"))
            origin = "manager_assigned"
        else:
            if not _active_field_session(db, user["id"]):
                flash("Start Field Work before creating an independent visit.", "error")
                return redirect(url_for("field_visits"))
            assigned_to = user["id"]
            origin = "executive_initiated"

        lead_id = f.get("lead_id", "").strip() or None
        company_id = f.get("company_id", "").strip() or None
        lead = db.execute("SELECT * FROM leads WHERE id=?", (lead_id,)).fetchone() if lead_id else None
        if lead_id and not lead:
            flash("The selected lead was not found.", "error")
            return redirect(url_for("field_visit_new"))
        if lead and lead["company_id"]:
            company_id = lead["company_id"]
        company = db.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone() if company_id else None
        if not lead and not company:
            flash("Select an existing lead or company.", "error")
            return redirect(url_for("field_visit_new"))

        scheduled_date = f.get("scheduled_date", "").strip() or today_str()
        scheduled_time = f.get("scheduled_time", "").strip() or datetime.now(APP_TIMEZONE).strftime("%H:%M")
        address = f.get("location_address", "").strip()
        purpose = f.get("purpose", "").strip()
        if not address or not purpose:
            flash("Visit location and purpose are required.", "error")
            return redirect(url_for("field_visit_new", lead_id=lead_id or "", company_id=company_id or ""))

        cur = db.execute(
            """INSERT INTO field_visits
               (lead_id,company_id,assigned_to,created_by,origin,scheduled_date,scheduled_time,
                location_address,purpose,contact_person,status,created_at,updated_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (lead_id, company_id, assigned_to, user["id"], origin, scheduled_date, scheduled_time,
             address, purpose, f.get("contact_person", "").strip(), "scheduled", now_str(), now_str()),
        )
        db.commit()
        visit = _get_field_visit(db, cur.lastrowid)
        _log_field_visit_event(visit, "Field Visit Created", f"{origin.replace('_', ' ').title()} · {scheduled_date} {scheduled_time}")
        flash("Field visit created.", "success")
        return redirect(url_for("field_visit_detail", visit_id=visit["id"]))

    executives = db.execute("SELECT id,name FROM users WHERE role='field_sales' ORDER BY name").fetchall()
    leads = db.execute(
        """SELECT leads.id,leads.name,leads.phone,leads.company_id,companies.name AS company_name,
                  companies.address,companies.billing_address
           FROM leads LEFT JOIN companies ON companies.id=leads.company_id
           ORDER BY leads.created_at DESC"""
    ).fetchall()
    companies = db.execute(
        "SELECT id,name,address,billing_address,city,state FROM companies ORDER BY name"
    ).fetchall()
    return render_template(
        "field_visit_form.html", executives=executives, leads=leads, companies=companies,
        selected_lead=request.args.get("lead_id", ""), selected_company=request.args.get("company_id", ""),
        now_time=datetime.now(APP_TIMEZONE).strftime("%H:%M"),
    )


@app.route("/field-visits/<int:visit_id>")
@field_sales_or_manager_required
def field_visit_detail(visit_id):
    db = get_db()
    user = current_user()
    visit = _get_field_visit(db, visit_id)
    if not _can_access_field_visit(user, visit):
        flash("Visit not found or not available to your account.", "error")
        return redirect(url_for("field_visits"))
    photos = db.execute(
        "SELECT * FROM field_visit_photos WHERE visit_id=? ORDER BY id", (visit_id,)
    ).fetchall()
    work_session = db.execute(
        "SELECT * FROM field_work_sessions WHERE id=?", (visit["session_id"],)
    ).fetchone() if visit["session_id"] else None
    active_session = _active_field_session(db, user["id"]) if user["role"] == "field_sales" else None
    return render_template(
        "field_visit_detail.html", visit=visit, photos=photos, work_session=work_session,
        active_session=active_session, outcomes=VISIT_OUTCOMES, no_card_reasons=NO_CARD_REASONS,
        session_duration=_session_duration,
    )


@app.route("/field-visits/<int:visit_id>/check-in", methods=["POST"])
@field_sales_or_manager_required
def field_visit_checkin(visit_id):
    db = get_db()
    user = current_user()
    visit = _get_field_visit(db, visit_id)
    if not visit or user["role"] != "field_sales" or visit["assigned_to"] != user["id"]:
        flash("Only the assigned executive can check in.", "error")
        return redirect(url_for("field_visits"))
    if visit["status"] != "scheduled":
        flash("This visit cannot be checked in again.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))
    work_session = _active_field_session(db, user["id"])
    if not work_session:
        flash("Start Field Work before checking in.", "error")
        return redirect(url_for("field_visits"))

    location_status = request.form.get("location_status", "unavailable")
    allowed_statuses = {"captured", "denied", "unavailable", "unsupported", "error"}
    if location_status not in allowed_statuses:
        location_status = "error"
    def optional_float(name):
        try:
            return float(request.form.get(name))
        except (TypeError, ValueError):
            return None
    lat = optional_float("latitude")
    lng = optional_float("longitude")
    accuracy = optional_float("accuracy")
    coordinates_valid = (
        lat is not None and lng is not None and -90 <= lat <= 90 and -180 <= lng <= 180
        and accuracy is not None and 0 <= accuracy <= 100000
    )
    if location_status == "captured" and not coordinates_valid:
        location_status = "error"
        lat = lng = accuracy = None
    elif location_status != "captured":
        lat = lng = accuracy = None

    db.execute(
        """UPDATE field_visits SET status='checked_in',session_id=?,checkin_at=?,checkin_lat=?,checkin_lng=?,
           checkin_accuracy=?,location_status=?,updated_at=? WHERE id=?""",
        (work_session["id"], now_str(), lat, lng, accuracy, location_status, now_str(), visit_id),
    )
    db.commit()
    visit = _get_field_visit(db, visit_id)
    _log_field_visit_event(visit, "Field Visit Check-in", f"Location: {location_status}")
    flash("Checked in. Complete the visit before ending field work.", "success")
    return redirect(url_for("field_visit_detail", visit_id=visit_id))


@app.route("/field-visits/<int:visit_id>/complete", methods=["POST"])
@field_sales_or_manager_required
def field_visit_complete(visit_id):
    db = get_db()
    user = current_user()
    visit = _get_field_visit(db, visit_id)
    if not visit or user["role"] != "field_sales" or visit["assigned_to"] != user["id"]:
        flash("Only the assigned executive can complete this visit.", "error")
        return redirect(url_for("field_visits"))
    if visit["status"] not in {"checked_in", "needs_correction"}:
        flash("Check in before completing this visit.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))
    if visit["status"] == "checked_in":
        active_session = _active_field_session(db, user["id"])
        if not active_session or active_session["id"] != visit["session_id"]:
            flash("The field-work session used for this check-in is no longer active.", "error")
            return redirect(url_for("field_visit_detail", visit_id=visit_id))

    f = request.form
    people_met = f.get("people_met", "").strip()
    discussion = f.get("discussion", "").strip()
    outcome = f.get("outcome", "").strip()
    contact_name = f.get("contact_name", "").strip()
    contact_designation = f.get("contact_designation", "").strip()
    contact_phone = f.get("contact_phone", "").strip()
    contact_email = f.get("contact_email", "").strip()
    if not people_met or not discussion or not outcome or not contact_name:
        flash("Contact name, people met, discussion, and visit outcome are required.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))

    office_data = f.get("office_photo_data", "")
    card_data = f.get("card_photo_data", "")
    has_office = bool(office_data) or _photo_exists(db, visit_id, "office")
    has_card = bool(card_data) or _photo_exists(db, visit_id, "visiting_card")
    no_card_reason = f.get("no_card_reason", "").strip()
    no_card_explanation = f.get("no_card_explanation", "").strip()
    if not has_office:
        flash("Capture an office photo with the camera before completing the visit.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))
    if not has_card and (no_card_reason not in NO_CARD_REASONS or not no_card_explanation):
        flash("Capture the visiting card, or choose a reason and enter an explanation.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))

    try:
        if office_data:
            _save_camera_photo(db, visit_id, "office", office_data)
        if card_data:
            _save_camera_photo(db, visit_id, "visiting_card", card_data)
            no_card_reason = ""
            no_card_explanation = ""
    except ValueError as exc:
        db.rollback()
        flash(str(exc), "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))

    completed_at = visit["completed_at"] or now_str()
    follow_up_date = f.get("follow_up_date", "").strip()
    db.execute(
        """UPDATE field_visits SET status='review_pending',completed_at=?,contact_person=?,contact_designation=?,contact_phone=?,contact_email=?,
           people_met=?,discussion=?,outcome=?,notes=?,
           next_action=?,follow_up_date=?,no_card_reason=?,no_card_explanation=?,submitted_at=?,review_status='pending',
           reviewed_by=NULL,reviewed_at=NULL,updated_at=? WHERE id=?""",
        (completed_at, contact_name, contact_designation, contact_phone, contact_email,
         people_met, discussion, outcome, f.get("notes", "").strip(),
         f.get("next_action", "").strip(), follow_up_date, no_card_reason, no_card_explanation,
         now_str(), now_str(), visit_id),
    )
    if visit["lead_id"]:
        db.execute(
            """UPDATE leads SET name=?,designation=?,phone=?,email=?,
               follow_up_date=CASE WHEN ?!='' THEN ? ELSE follow_up_date END,updated_at=? WHERE id=?""",
            (contact_name, contact_designation, contact_phone, contact_email,
             follow_up_date, follow_up_date, now_str(), visit["lead_id"]),
        )
        if visit["lead_contact_id"]:
            db.execute(
                "UPDATE contacts SET name=?,designation=?,phone=?,email=?,updated_at=? WHERE id=?",
                (contact_name, contact_designation, contact_phone, contact_email, now_str(), visit["lead_contact_id"]),
            )
    elif visit["company_id"]:
        existing_contact = db.execute(
            "SELECT id FROM contacts WHERE company_id=? AND LOWER(name)=LOWER(?) LIMIT 1",
            (visit["company_id"], contact_name),
        ).fetchone()
        if existing_contact:
            db.execute(
                """UPDATE contacts SET designation=?,
                   phone=CASE WHEN ?!='' THEN ? ELSE phone END,
                   email=CASE WHEN ?!='' THEN ? ELSE email END,updated_at=? WHERE id=?""",
                (contact_designation, contact_phone, contact_phone, contact_email, contact_email, now_str(), existing_contact["id"]),
            )
        else:
            db.execute(
                """INSERT INTO contacts (company_id,name,designation,phone,email,is_primary,created_at,updated_at)
                   VALUES (?,?,?,?,?,0,?,?)""",
                (visit["company_id"], contact_name, contact_designation, contact_phone, contact_email, now_str(), now_str()),
            )
    db.commit()
    visit = _get_field_visit(db, visit_id)
    _log_field_visit_event(visit, "Field Visit Submitted", f"Outcome: {outcome}" + (f" · Next: {f.get('next_action', '').strip()}" if f.get("next_action", "").strip() else ""))
    flash("Visit submitted for manager review.", "success")
    return redirect(url_for("field_visit_detail", visit_id=visit_id))


@app.route("/field-visits/<int:visit_id>/review", methods=["POST"])
@manager_required
def field_visit_review(visit_id):
    db = get_db()
    visit = _get_field_visit(db, visit_id)
    if not visit or visit["review_status"] != "pending":
        flash("This visit is not waiting for review.", "error")
        return redirect(url_for("field_visits"))
    action = request.form.get("action")
    comment = request.form.get("review_comment", "").strip()
    if action == "return" and not comment:
        flash("Enter a correction comment before returning the visit.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))
    if action not in {"accept", "return"}:
        flash("Choose Accept or Return for correction.", "error")
        return redirect(url_for("field_visit_detail", visit_id=visit_id))
    status = "completed" if action == "accept" else "needs_correction"
    review_status = "accepted" if action == "accept" else "returned"
    db.execute(
        """UPDATE field_visits SET status=?,review_status=?,reviewed_by=?,reviewed_at=?,review_comment=?,updated_at=?
           WHERE id=?""",
        (status, review_status, current_user()["id"], now_str(), comment, now_str(), visit_id),
    )
    db.commit()
    visit = _get_field_visit(db, visit_id)
    event = "Field Visit Accepted" if action == "accept" else "Field Visit Returned"
    _log_field_visit_event(visit, event, comment)
    flash("Visit accepted." if action == "accept" else "Visit returned for correction.", "success")
    return redirect(url_for("field_visit_detail", visit_id=visit_id))


@app.route("/field-visit-photos/<int:photo_id>")
@field_sales_or_manager_required
def field_visit_photo(photo_id):
    db = get_db()
    row = db.execute("SELECT * FROM field_visit_photos WHERE id=?", (photo_id,)).fetchone()
    visit = _get_field_visit(db, row["visit_id"]) if row else None
    if not row or not _can_access_field_visit(current_user(), visit):
        return "Not found", 404
    path = os.path.join(_field_photo_dir(), row["stored_name"])
    if not os.path.isfile(path):
        return "Image not found", 404
    return send_file(path, mimetype=row["mime_type"], as_attachment=False, download_name=f"{row['photo_type']}.jpg")


# ---------- Invoices (manager only) ----------

def _invoice_settings(db):
    return db.execute("SELECT * FROM company_settings WHERE id=1").fetchone()


def _next_invoice_no(db, prefix):
    today = date.today()
    fy_start = today.year if today.month >= 4 else today.year - 1
    fy = f"{str(fy_start)[-2:]}-{str(fy_start + 1)[-2:]}"
    stem = f"{prefix or 'QL'}/{fy}/"
    rows = db.execute("SELECT invoice_no FROM invoices WHERE invoice_no LIKE ?", (stem + '%',)).fetchall()
    nums = []
    for r in rows:
        try: nums.append(int(r['invoice_no'].rsplit('/',1)[1]))
        except Exception: pass
    return f"{stem}{(max(nums, default=0)+1):05d}"


@app.route('/invoices')
@manager_required
def invoices_list():
    db=get_db()
    invoices=db.execute("SELECT * FROM invoices ORDER BY id DESC").fetchall()
    return render_template('invoices_list.html', invoices=invoices)


@app.route('/invoices/new', methods=['GET','POST'])
@manager_required
def invoice_new():
    db=get_db(); settings=_invoice_settings(db)
    if request.method == 'GET':
        companies=db.execute("SELECT * FROM companies ORDER BY name").fetchall()
        return render_template('invoice_form.html', companies=companies, settings=settings,
                               next_invoice_no=_next_invoice_no(db, settings['invoice_prefix']))
    f=request.form
    names=f.getlist('item_name[]'); codes=f.getlist('item_code[]'); qtys=f.getlist('item_qty[]'); rates=f.getlist('item_rate[]'); taxes=f.getlist('item_tax[]')
    items=[]
    for i,name in enumerate(names):
        if not name.strip(): continue
        try: qty=float(qtys[i]); rate=float(rates[i]); tax=float(taxes[i])
        except (ValueError,IndexError):
            flash('Every invoice item needs valid quantity, rate and tax values.','error'); return redirect(url_for('invoice_new'))
        items.append((name.strip(), codes[i].strip() if i < len(codes) else '', qty, rate, tax))
    if not items or not f.get('bill_name','').strip():
        flash('Billing company and at least one item are required.','error'); return redirect(url_for('invoice_new'))
    invoice_no=f.get('invoice_no','').strip() or _next_invoice_no(db,settings['invoice_prefix'])

    # GST is determined from the shipping state.
    # Telangana = intra-State (CGST + SGST); any other state = inter-State (IGST).
    ship_state=f.get('ship_state','').strip()
    is_interstate=bool(ship_state) and ship_state.casefold() != 'telangana'
    if is_interstate:
        tax_mode='INTER'
        cgst=0.0
        sgst=0.0
        igst=float(f.get('igst_pct') or 18)
    else:
        tax_mode='INTRA'
        cgst=float(f.get('cgst_pct') or 9)
        sgst=float(f.get('sgst_pct') or 9)
        igst=0.0

    received=float(f.get('received_amount') or 0)
    total_tax_pct = igst if is_interstate else (cgst + sgst)
    total=sum(q*r*(1+total_tax_pct/100) for _,_,q,r,_ in items)

    try:
        cur=db.execute("""INSERT INTO invoices (invoice_no,invoice_date,company_id,bill_name,bill_address,bill_gstin,bill_pan,bill_website,
            ship_name,ship_address,ship_gstin,ship_pan,ship_website,ship_state,notes,terms,cgst_pct,sgst_pct,igst_pct,tax_mode,received_amount,total_amount,created_by,created_at,
            bill_contact_name,bill_contact_email,bill_contact_phone)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(
            invoice_no,f.get('invoice_date','').strip(),f.get('company_id') or None,f.get('bill_name','').strip(),f.get('bill_address','').strip(),
            f.get('bill_gstin','').strip().upper(),f.get('bill_pan','').strip().upper(),f.get('bill_website','').strip(),
            f.get('ship_name','').strip(),f.get('ship_address','').strip(),f.get('ship_gstin','').strip().upper(),f.get('ship_pan','').strip().upper(),f.get('ship_website','').strip(),
            ship_state,f.get('notes','').strip(),f.get('terms','').strip(),cgst,sgst,igst,tax_mode,received,total,current_user()['id'],now_str(),
            f.get('bill_contact_name','').strip(), f.get('bill_contact_email','').strip(), f.get('bill_contact_phone','').strip()))
        iid=cur.lastrowid
        db.executemany("INSERT INTO invoice_items (invoice_id,item_name,code,qty,rate,tax_pct) VALUES (?,?,?,?,?,?)",[(iid,*x) for x in items]); db.commit()
    except sqlite3.IntegrityError:
        flash('That invoice number already exists. Please use a different number.','error'); return redirect(url_for('invoice_new'))
    log_activity('Invoice Created', 'invoice', iid, invoice_no, f.get('bill_name','').strip())
    flash(f'Invoice {invoice_no} created.','success'); return redirect(url_for('invoice_detail', invoice_id=iid))


@app.route('/invoices/<int:invoice_id>')
@manager_required
def invoice_detail(invoice_id):
    db=get_db(); inv=db.execute("SELECT * FROM invoices WHERE id=?",(invoice_id,)).fetchone()
    if not inv: flash('Invoice not found.','error'); return redirect(url_for('invoices_list'))
    items=db.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id",(invoice_id,)).fetchall()
    return render_template('invoice_detail.html', invoice=inv, items=items)


@app.route('/invoices/<int:invoice_id>/pdf')
@manager_required
def invoice_pdf(invoice_id):
    db=get_db(); inv=db.execute("SELECT * FROM invoices WHERE id=?",(invoice_id,)).fetchone(); settings=_invoice_settings(db)
    if not inv: flash('Invoice not found.','error'); return redirect(url_for('invoices_list'))
    rows=db.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id",(invoice_id,)).fetchall()
    company={'name':settings['name'],'address_lines':settings['address_lines'],'website':settings['website'],'gstin':settings['gstin'],'pan':settings['pan']}
    bill={'name':inv['bill_name'],'address_lines':inv['bill_address'],'gstin':inv['bill_gstin'],'pan':inv['bill_pan'],'website':inv['bill_website']}
    ship={'name':inv['ship_name'] or inv['bill_name'],'address_lines':inv['ship_address'] or inv['bill_address'],'gstin':inv['ship_gstin'] or inv['bill_gstin'],'pan':inv['ship_pan'] or inv['bill_pan'],'website':inv['ship_website'] or inv['bill_website'],'state':inv['ship_state'] if 'ship_state' in inv.keys() else ''}
    items=[{'name':r['item_name'],'code':r['code'],'qty':r['qty'],'rate':r['rate'],'tax_pct':r['tax_pct']} for r in rows]
    bank={'name':settings['bank_name'],'ifsc':settings['ifsc'],'account_no':settings['account_no'],'bank':settings['bank_branch'],'upi_id':settings['upi_id']}
    pdf=build_invoice_pdf(company,bill,ship,inv['invoice_no'],inv['invoice_date'],items,inv['notes'],(inv['terms'] or '').splitlines(),bank,inv['cgst_pct'],inv['sgst_pct'],inv['received_amount'],inv['igst_pct'] if 'igst_pct' in inv.keys() else 0,inv['tax_mode'] if 'tax_mode' in inv.keys() else 'INTRA',inv['ship_state'] if 'ship_state' in inv.keys() else '')
    safe=inv['invoice_no'].replace('/','-')
    log_activity('Invoice Downloaded', 'invoice', invoice_id, inv['invoice_no'])
    return send_file(pdf,mimetype='application/pdf',as_attachment=True,download_name=f'invoice_{safe}.pdf')


@app.route('/invoices/<int:invoice_id>/delete', methods=['POST'])
@manager_required
def delete_invoice(invoice_id):
    db=get_db()
    inv=db.execute("SELECT * FROM invoices WHERE id=?",(invoice_id,)).fetchone()
    if not inv:
        flash('Invoice not found.','error'); return redirect(url_for('invoices_list'))
    db.execute("DELETE FROM invoice_items WHERE invoice_id=?", (invoice_id,))
    db.execute("DELETE FROM invoices WHERE id=?", (invoice_id,))
    db.commit()
    log_activity('Invoice Deleted', 'invoice', invoice_id, inv['invoice_no'])
    flash(f"Invoice {inv['invoice_no']} deleted.", 'success')
    return redirect(url_for('invoices_list'))


@app.route('/settings/company', methods=['GET','POST'])
@manager_required
def company_settings():
    db=get_db()
    if request.method=='POST':
        f=request.form
        db.execute("""UPDATE company_settings SET name=?,address_lines=?,website=?,gstin=?,pan=?,bank_name=?,ifsc=?,account_no=?,bank_branch=?,upi_id=?,default_terms=?,default_notes=?,invoice_prefix=? WHERE id=1""",(
            f.get('name','').strip(),f.get('address_lines','').strip(),f.get('website','').strip(),f.get('gstin','').strip().upper(),f.get('pan','').strip().upper(),
            f.get('bank_name','').strip(),f.get('ifsc','').strip().upper(),f.get('account_no','').strip(),f.get('bank_branch','').strip(),f.get('upi_id','').strip(),
            f.get('default_terms','').strip(),f.get('default_notes','').strip(),f.get('invoice_prefix','').strip().upper() or 'QL')); db.commit()
        flash('Company invoice settings saved.','success'); return redirect(url_for('company_settings'))
    return render_template('company_settings.html', settings=_invoice_settings(db))


@app.context_processor
def inject_globals():
    return {"today": today_str(), "current_user": current_user()}


@app.template_filter("slug")
def slug_filter(value):
    return (value or "").replace(" ", "-")


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)
else:
    init_db()
